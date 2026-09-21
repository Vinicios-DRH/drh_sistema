"""
Geração da Nota para BG da Junta Médica (JOIS/CBMAM).

Construído com python-docx puro — sem Jinja, sem docxtpl. O documento inteiro
(seções, tabelas, errata, assinaturas) é montado em código a partir do casco
`modelo_nota_bg.docx`, que só carrega o timbre oficial (cabeçalho/rodapé,
margens, fontes) com o corpo vazio.

Essa escolha existe por causa de uma dor real: a versão anterior usava um
.docx com tags Jinja escritas à mão dentro das tabelas do Word, e o Word quebra
essas tags em vários "runs" XML sem avisar (basta clicar no meio do texto e
digitar) — o que já corrompeu a geração da nota mais de uma vez. Gerando tudo
via código, esse tipo de corrupção silenciosa deixa de ser possível.
"""
from __future__ import annotations

from itertools import groupby
from pathlib import Path
from datetime import date, timedelta
from typing import Optional
from zoneinfo import ZoneInfo

import docx
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from sqlalchemy.orm import joinedload

from src import database
from src.models import JuntaFechamentoBg, Licencas, LicencaRestricao, Militar
from src.services.junta_medica import ordem_hierarquica_militar

MANAUS_TZ = ZoneInfo("America/Manaus")

TEMPLATE_PATH = Path("src/template/modelo_nota_bg.docx")
OUTPUT_DIR = Path("src/static/junta_bg")

FONTE_PADRAO = "Arial"
TAMANHO_PADRAO = Pt(11)

MESES_PT = [
    "janeiro", "fevereiro", "março", "abril", "maio", "junho",
    "julho", "agosto", "setembro", "outubro", "novembro", "dezembro",
]

# Corpo permanente da JOIS que assina a nota. Só muda quando a composição da
# junta muda de fato — editar aqui é o único lugar necessário.
MEMBROS_JOIS = [
    {"nome": "SILVANA DE LIMA E SILVA", "identificacao": "CAP QCOBM Méd.",
     "cargo": "Presidente da JOIS/CBMAM", "crm": "CRM-AM 3492"},
    {"nome": "KARLA GODINHO DO CARMO SILVA", "identificacao": "CAP QCOBM Méd.",
     "cargo": "Membro da JOIS/CBMAM", "crm": "CRM-AM 5848"},
    {"nome": "NAYARA DE ALENCAR DIAS", "identificacao": "1º TEN QCOBM Méd.",
     "cargo": "Membro da JOIS/CBMAM", "crm": "CRM-AM 7521"},
]


# ---------------------------------------------------------------------------
# Formatação de texto
# ---------------------------------------------------------------------------

def fmt_data_ponto(dt) -> str:
    """JOIS usa data com pontos (25.07.2026), não barras."""
    if not dt:
        return ""
    return dt.strftime("%d.%m.%Y")


def data_por_extenso(dt) -> str:
    if not dt:
        return ""
    return f"{dt.day:02d} de {MESES_PT[dt.month - 1]} de {dt.year}"


def _data_por_extenso_maiuscula(dt) -> str:
    """Como `data_por_extenso`, mas em caixa alta — pro meio de frases que
    já são todas em maiúsculo (AO/ISO)."""
    return data_por_extenso(dt).upper()


def safe_getattr(obj, attr, default=""):
    try:
        value = getattr(obj, attr, default)
        return value if value is not None else default
    except Exception:
        return default


def numero_por_extenso(n: int) -> str:
    unidades = {
        0: "zero", 1: "um", 2: "dois", 3: "três", 4: "quatro", 5: "cinco",
        6: "seis", 7: "sete", 8: "oito", 9: "nove", 10: "dez",
        11: "onze", 12: "doze", 13: "treze", 14: "quatorze", 15: "quinze",
        16: "dezesseis", 17: "dezessete", 18: "dezoito", 19: "dezenove"
    }
    dezenas = {
        20: "vinte", 30: "trinta", 40: "quarenta", 50: "cinquenta",
        60: "sessenta", 70: "setenta", 80: "oitenta", 90: "noventa"
    }
    centenas = {
        100: "cem", 200: "duzentos", 300: "trezentos", 400: "quatrocentos",
        500: "quinhentos", 600: "seiscentos", 700: "setecentos",
        800: "oitocentos", 900: "novecentos"
    }

    if n < 0:
        return str(n)
    if n in unidades:
        return unidades[n]
    if n < 100:
        dez = (n // 10) * 10
        resto = n % 10
        return dezenas[dez] if resto == 0 else f"{dezenas[dez]} e {unidades[resto]}"
    if n == 100:
        return "cem"
    if n < 1000:
        cent = (n // 100) * 100
        resto = n % 100
        cent_texto = "cento" if cent == 100 else centenas[cent]
        return cent_texto if resto == 0 else f"{cent_texto} e {numero_por_extenso(resto)}"
    return str(n)


def _lista_com_e(itens) -> str:
    """['A', 'B', 'C'] -> 'A, B E C'. Usado nas listas de restrição."""
    itens = list(itens)
    if not itens:
        return ""
    if len(itens) == 1:
        return itens[0]
    return ", ".join(itens[:-1]) + " E " + itens[-1]


def texto_restricoes_bg(lic: Licencas) -> str:
    """
    Linha de restrições no padrão real da JOIS:
    "RECOMENDAÇÕES: RESTRIÇÃO PARA TAF, TFM E FORMATURA."
    """
    nomes = sorted(
        {v.tipo.nome for v in getattr(lic, "restricoes", []) if v.tipo},
        key=str.lower
    )
    if not nomes:
        return ""
    return f"RECOMENDAÇÕES: RESTRIÇÃO PARA {_lista_com_e(nomes)}."


def _mascarar_rg(rg: Optional[str]) -> str:
    """
    IDT da nota: o RG nunca aparece inteiro — só os 2 últimos dígitos, com
    "**" no lugar do resto. É o padrão oficial da JOIS (ver modelo de
    referência: toda identificação sai como "**00", "**19", "**64"...).
    """
    rg = (rg or "").strip()
    return f"**{rg[-2:]}" if rg else "**"


def _identidade_militar(lic: Licencas):
    militar = lic.militar
    posto_grad = safe_getattr(safe_getattr(militar, "posto_grad", None), "sigla", "")
    quadro = safe_getattr(safe_getattr(militar, "quadro", None), "quadro", "")
    nome = safe_getattr(militar, "nome_completo", "")
    rg = safe_getattr(militar, "rg", "")

    return {
        "posto_grad_quadro": f"{posto_grad} {quadro}".strip(),
        "nome_completo": nome,
        "idt": f"{_mascarar_rg(rg)}\nCBMAM",
    }


def _precisa_reavaliar(lic: Licencas) -> bool:
    """
    Decide qual texto a coluna de situação leva. Vale pra LTS (a pedido do
    usuário) e, seguindo o mesmo padrão real da JOIS, também pra
    APTO_RESTR/APTO_RECOM — o modelo oficial mostra exatamente essa mesma
    marcação nesses pareceres. LTSPF/LM não têm essa decisão: encerrado o
    prazo, sempre voltam a ser aptos sozinhos.
    """
    if lic.tipo_licenca in ("LTS", "APTO_RESTR", "APTO_RECOM"):
        return bool(lic.reavaliar_ao_termino)
    return False


def _texto_situacao(lic: Licencas) -> str:
    if _precisa_reavaliar(lic):
        return "REAVALIAR AO TÉRMINO"
    pronto_em = lic.data_fim + timedelta(days=1)
    return f"PRONTO PARA SV\n{fmt_data_ponto(pronto_em)}"


def _texto_periodo(lic: Licencas) -> str:
    unidade = "DIA" if lic.qtd_dias == 1 else "DIAS"
    extenso = numero_por_extenso(lic.qtd_dias).upper()
    return (
        f"POR: {lic.qtd_dias:02d} ({extenso}) {unidade} A/C DE "
        f"{fmt_data_ponto(lic.data_inicio)}\nTÉRMINO: {fmt_data_ponto(lic.data_fim)}"
    )


def _linhas_nome(lic: Licencas, com_periodo: bool) -> list[str]:
    linhas = [_identidade_militar(lic)["nome_completo"]]

    if com_periodo:
        linhas.append(_texto_periodo(lic))

    restricoes = texto_restricoes_bg(lic)
    if restricoes:
        linhas.append(restricoes)

    if lic.observacao:
        linhas.append(lic.observacao.strip())

    return linhas


# ---------------------------------------------------------------------------
# Construção do documento — helpers de baixo nível
# ---------------------------------------------------------------------------

def _aplicar_fonte(run, negrito=False, sublinhado=False, tamanho=None):
    run.font.name = FONTE_PADRAO
    run.font.size = tamanho or TAMANHO_PADRAO
    run.bold = negrito
    run.underline = sublinhado


def _paragrafo(doc, texto="", negrito=False, sublinhado=False,
              alinhamento=WD_ALIGN_PARAGRAPH.JUSTIFY, espaco_depois=6):
    p = doc.add_paragraph()
    p.alignment = alinhamento
    p.paragraph_format.space_after = Pt(espaco_depois)

    if texto:
        run = p.add_run(texto)
        _aplicar_fonte(run, negrito=negrito, sublinhado=sublinhado)

    return p


def _centralizar_celula(cell):
    """Conteúdo sempre centralizado (horizontal e vertical) — é o padrão
    de toda tabela da nota, cabeçalho ou dado."""
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    for p in cell.paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _paragrafo_multilinha(cell, linhas: list[str], negrito_primeira=False):
    """Escreve várias linhas numa célula de tabela, cada uma seu parágrafo."""
    cell.text = ""
    primeiro = True

    for linha in linhas:
        p = cell.paragraphs[0] if primeiro else cell.add_paragraph()
        run = p.add_run(linha)
        _aplicar_fonte(run, negrito=(primeiro and negrito_primeira))
        primeiro = False

    _centralizar_celula(cell)


def _tabela(doc, cabecalho: list[str]):
    tabela = doc.add_table(rows=1, cols=len(cabecalho))
    tabela.style = "Table Grid"

    for idx, texto in enumerate(cabecalho):
        cell = tabela.rows[0].cells[idx]
        cell.text = ""
        run = cell.paragraphs[0].add_run(texto)
        _aplicar_fonte(run, negrito=True)
        _centralizar_celula(cell)

    return tabela


def _linha_tabela(tabela, valores: list):
    """
    `valores` aceita string simples (uma linha) ou lista de strings
    (multi-linha, uma por parágrafo) em cada coluna.
    """
    row = tabela.add_row()
    for idx, valor in enumerate(valores):
        cell = row.cells[idx]
        if isinstance(valor, (list, tuple)):
            _paragrafo_multilinha(cell, list(valor))
        else:
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(valor))
            _aplicar_fonte(run)
            _centralizar_celula(cell)
    return row


# ---------------------------------------------------------------------------
# Seções de licença/situação (LTS, LTSPF, LM, APTO_RECOM, APTO_RESTR)
# ---------------------------------------------------------------------------

SECOES_PERIODO = [
    ("LTS", "INCAPAZES TEMPORARIAMENTE PARA O SERVIÇO DO CBMAM: OS BMS "
            "DISCRIMINADOS ABAIXO ESTÃO APTOS A RESPONDER PROCESSOS "
            "ADMINISTRATIVOS E OUTROS DE NATUREZA JUDICIAL."),
    ("LTSPF", "LICENÇA PARA TRATAMENTO DE SAÚDE DE PESSOA DA FAMÍLIA (LTSPF): "
              "OS BMS DISCRIMINADOS ABAIXO ESTÃO EM GOZO DA LICENÇA."),
    ("LM", "LICENÇA MATERNIDADE: AS BMS DISCRIMINADAS ABAIXO ESTÃO EM GOZO "
           "DA LICENÇA."),
    ("APTO_RECOM", "APTOS COM RECOMENDAÇÕES PARA O SERVIÇO DO CBMAM: OS BMS "
                   "DISCRIMINADOS ABAIXO ESTÃO APTOS PARA O SERVIÇO "
                   "ADMINISTRATIVO, A RESPONDER PROCESSOS ADMINISTRATIVOS E "
                   "OUTROS DE NATUREZA JUDICIAL."),
    ("APTO_RESTR", "APTOS COM RESTRIÇÕES PARA O SERVIÇO DO CBMAM: OS BMS "
                   "DISCRIMINADOS ABAIXO ESTÃO APTOS PARA O SERVIÇO DO CBMAM, "
                   "COM AS RESTRIÇÕES INDICADAS."),
]


def _secao_periodo(doc, tipo: str, cabecalho_texto: str, itens: list[Licencas]):
    if not itens:
        return

    _paragrafo(doc, cabecalho_texto, espaco_depois=4)

    tabela = _tabela(doc, ["POSTO/GRAD", "NOME", "SITUAÇÃO AO TÉRMINO", "IDT"])
    for lic in itens:
        ident = _identidade_militar(lic)
        _linha_tabela(tabela, [
            ident["posto_grad_quadro"],
            _linhas_nome(lic, com_periodo=True),
            _texto_situacao(lic).split("\n"),
            ident["idt"].split("\n"),
        ])

    _paragrafo(doc, "", espaco_depois=10)


# ---------------------------------------------------------------------------
# Seções pontuais (CURSO / TAF / PROMOÇÃO) — grupos dinâmicos
# ---------------------------------------------------------------------------

def _titulo_curso(resultado: str, curso_nome: str, detalhe: Optional[str],
                  numero_bg: Optional[str] = None) -> str:
    if resultado == "CURSO_APTO":
        base = f"APTOS PARA CURSO DE {curso_nome}"
    elif resultado == "CURSO_INAPTO":
        base = f"INAPTOS PARA CURSO DE {curso_nome}"
    elif resultado == "CURSO_REGIME_ESPECIAL":
        base = f"EM REGIME ESPECIAL PARA CURSO DE {curso_nome}"
    elif resultado == "CURSO_OUTRO":
        rotulo = (detalhe or "OUTRO RESULTADO").strip().upper()
        base = f"{rotulo} — CURSO DE {curso_nome}"
    else:
        base = f"CURSO DE {curso_nome}"

    # Número do BG interno que o operador informou na inspeção — é a
    # referência que ele digitou na tela, precisa continuar aparecendo aqui.
    return _com_numero_bg(base, numero_bg)


def _com_numero_bg(base_sem_ponto: str, numero_bg: Optional[str]) -> str:
    if numero_bg:
        return f"{base_sem_ponto}, CONFORME BG Nº {numero_bg}."
    return f"{base_sem_ponto}."


def _titulo_taf(resultado: str, numero_bg: Optional[str] = None) -> str:
    base = {
        "TAF_APTO": "APTOS PARA FINS DE TESTE DE APTIDÃO FÍSICA (TAF)",
        "TAF_ALTERNATIVO": "APTOS PARA TAF ALTERNATIVO",
        "TAF_INAPTO": "INAPTOS, PARA FINS DE TESTE DE APTIDÃO FÍSICA (TAF)",
    }.get(resultado, "RESULTADO DO TAF")
    return _com_numero_bg(base, numero_bg)


def _titulo_promocao(resultado: str, numero_bg: Optional[str] = None) -> str:
    base = {
        "PROMOCAO_APTO": "APTOS PARA FINS DE PROMOÇÃO",
        "PROMOCAO_INAPTO": "INAPTOS PARA FINS DE PROMOÇÃO",
    }.get(resultado, "RESULTADO PARA FINS DE PROMOÇÃO")
    return _com_numero_bg(base, numero_bg)


def _secao_grupo_simples(doc, titulo: str, itens: list[Licencas]):
    """Tabela de 2 colunas (POSTO/GRAD, NOME) — usada por CURSO/TAF/PROMOÇÃO."""
    _paragrafo(doc, titulo, espaco_depois=4)

    tabela = _tabela(doc, ["POSTO/GRAD", "NOME"])
    for lic in itens:
        ident = _identidade_militar(lic)
        linhas_nome = [ident["nome_completo"]]
        if lic.observacao:
            linhas_nome.append(lic.observacao.strip())

        _linha_tabela(tabela, [ident["posto_grad_quadro"], linhas_nome])

    _paragrafo(doc, "", espaco_depois=10)


def _secao_curso(doc, itens: list[Licencas]):
    chave = lambda lic: (lic.curso_nome or (lic.curso.nome if lic.curso else ""),
                         lic.status, lic.resultado_detalhe or "",
                         lic.numero_bg_curso or "")
    for (curso_nome, resultado, detalhe, numero_bg), grupo in groupby(
        sorted(itens, key=chave), key=chave
    ):
        titulo = _titulo_curso(resultado, curso_nome or "CURSO NÃO INFORMADO",
                               detalhe, numero_bg)
        _secao_grupo_simples(doc, titulo, list(grupo))


def _secao_taf(doc, itens: list[Licencas]):
    chave = lambda lic: (lic.status, lic.numero_bg_curso or "")
    for (resultado, numero_bg), grupo in groupby(sorted(itens, key=chave), key=chave):
        _secao_grupo_simples(doc, _titulo_taf(resultado, numero_bg), list(grupo))


def _secao_promocao(doc, itens: list[Licencas]):
    chave = lambda lic: (lic.status, lic.numero_bg_curso or "")
    for (resultado, numero_bg), grupo in groupby(sorted(itens, key=chave), key=chave):
        _secao_grupo_simples(doc, _titulo_promocao(resultado, numero_bg), list(grupo))


# ---------------------------------------------------------------------------
# AO (Exame de Controle de Atestado de Origem) e ISO (Parecer Técnico de
# Inquérito Sanitário de Origem) — pontuais, agrupados por portaria/data de
# publicação (é o que muda o cabeçalho da seção).
# ---------------------------------------------------------------------------

TEXTOS_RESULTADO_AO = {
    "AO_HA_RELACAO": (
        "HÁ RELAÇÃO DE CAUSA E EFEITO ENTRE O ACIDENTE SOFRIDO E AS "
        "CONDIÇÕES MÓRBIDAS ATUAIS EXPRESSAS PELOS DIAGNÓSTICOS."),
    "AO_NAO_HA_RELACAO": (
        "NÃO HÁ RELAÇÃO DE CAUSA E EFEITO ENTRE AS CONDIÇÕES INERENTES AO "
        "SERVIÇO E O ESTADO MÓRBIDO ATUAL EXPRESSO PELO DIAGNÓSTICO."),
    "AO_HA_VESTIGIOS":
        "HÁ VESTÍGIOS ANATÔMICOS OU FUNCIONAIS DO ACIDENTE SOFRIDO.",
    "AO_NAO_HA_VESTIGIOS":
        "NÃO HÁ VESTÍGIOS ANATÔMICOS OU FUNCIONAIS DO ACIDENTE SOFRIDO.",
}


def _texto_resultado_iso(lic: Licencas) -> str:
    """
    ISO_HA_RELACAO/ISO_NAO_HA_RELACAO tratam de um acidente específico — o
    texto cita a data dele. Os "_SERVICO" tratam de condições inerentes ao
    serviço em geral, sem data nenhuma.
    """
    data_ac = _data_por_extenso_maiuscula(lic.data_acidente)

    if lic.status == "ISO_HA_RELACAO":
        return (
            f"HÁ RELAÇÃO DE CAUSA E EFEITO ENTRE O ACIDENTE SOFRIDO NO DIA "
            f"{data_ac} E AS CONDIÇÕES MÓRBIDAS ATUAIS EXPRESSAS PELOS "
            f"DIAGNÓSTICOS."
        )
    if lic.status == "ISO_NAO_HA_RELACAO":
        return (
            f"NÃO HÁ RELAÇÃO DE CAUSA E EFEITO ENTRE O ACIDENTE SOFRIDO NO "
            f"DIA {data_ac} E AS CONDIÇÕES MÓRBIDAS ATUAIS EXPRESSAS PELOS "
            f"DIAGNÓSTICOS."
        )
    if lic.status == "ISO_HA_RELACAO_SERVICO":
        return (
            "HÁ RELAÇÃO DE CAUSA E EFEITO ENTRE AS CONDIÇÕES INERENTES AO "
            "SERVIÇO E O ESTADO MÓRBIDO ATUAL EXPRESSO PELOS DIAGNÓSTICOS."
        )
    if lic.status == "ISO_NAO_HA_RELACAO_SERVICO":
        return (
            "NÃO HÁ RELAÇÃO DE CAUSA E EFEITO ENTRE AS CONDIÇÕES INERENTES "
            "AO SERVIÇO E O ESTADO MÓRBIDO ATUAL EXPRESSO PELOS "
            "DIAGNÓSTICOS."
        )
    return ""


def _titulo_ao(portaria: str, data_publicacao) -> str:
    return (
        f"COMPARECEU A JOIS/CBMAM: FINS INSPEÇÃO DE SAÚDE PARA EXAME DE "
        f"CONTROLE DO ATESTADO DE ORIGEM, CONFORME PORTARIA Nº {portaria}, "
        f"DE {_data_por_extenso_maiuscula(data_publicacao)}, COM SEU "
        f"RESPECTIVO PARECER MÉDICO."
    )


def _titulo_iso(portaria: str, data_publicacao) -> str:
    return (
        f"COMPARECEU A JOIS/CBMAM: FINS INSPEÇÃO DE SAÚDE PARA PARECER "
        f"TÉCNICO DO INQUÉRITO SANITÁRIO DE ORIGEM - ISO, CONFORME PORTARIA "
        f"Nº {portaria}, DE {_data_por_extenso_maiuscula(data_publicacao)}, "
        f"COM SEU RESPECTIVO PARECER MÉDICO."
    )


def _secao_ao_iso(doc, tipo: str, itens: list[Licencas]):
    if not itens:
        return

    titulo_fn = _titulo_ao if tipo == "AO" else _titulo_iso
    chave = lambda lic: (lic.portaria or "", lic.data_publicacao or date.min)

    for (portaria, data_publicacao), grupo in groupby(
        sorted(itens, key=chave), key=chave
    ):
        _paragrafo(doc, titulo_fn(portaria, data_publicacao), espaco_depois=4)

        tabela = _tabela(doc, ["POSTO/GRAD", "NOME", "IDT"])
        for lic in grupo:
            ident = _identidade_militar(lic)
            texto_resultado = (
                TEXTOS_RESULTADO_AO.get(lic.status, "") if tipo == "AO"
                else _texto_resultado_iso(lic)
            )
            _linha_tabela(tabela, [
                ident["posto_grad_quadro"],
                [ident["nome_completo"], texto_resultado],
                ident["idt"].split("\n"),
            ])

        _paragrafo(doc, "", espaco_depois=10)


# ---------------------------------------------------------------------------
# AGREGADO e APTO (individuais, "a contar de" / "agregado a partir de")
# ---------------------------------------------------------------------------

def _secao_agrupada_por_data(doc, itens: list[Licencas], rotulo_singular: str):
    """
    AGREGADO e APTO viram um lançamento por data: agrupa quem tem a mesma
    data_inicio num único cabeçalho + tabela de 3 colunas (com IDT), igual ao
    "APTO AO SERVIÇO DO CBMAM A CONTAR DE DD.MM.YYYY." da JOIS.
    """
    chave = lambda lic: lic.data_inicio
    for data_inicio, grupo in groupby(sorted(itens, key=chave), key=chave):
        grupo = list(grupo)
        titulo = f"{rotulo_singular} A CONTAR DE {fmt_data_ponto(data_inicio)}."

        _paragrafo(doc, titulo, espaco_depois=4)
        tabela = _tabela(doc, ["POSTO/GRAD", "NOME", "IDT"])

        for lic in grupo:
            ident = _identidade_militar(lic)
            linhas_nome = [ident["nome_completo"]]
            if lic.observacao:
                linhas_nome.append(lic.observacao.strip())
            _linha_tabela(tabela, [
                ident["posto_grad_quadro"], linhas_nome, ident["idt"].split("\n")
            ])

        _paragrafo(doc, "", espaco_depois=10)


# ---------------------------------------------------------------------------
# Errata
# ---------------------------------------------------------------------------

def _secao_errata(doc, fechamento: JuntaFechamentoBg):
    if not fechamento.eh_errata:
        return

    original = fechamento.fechamento_original
    ano = fechamento.data_referencia.year

    _paragrafo(
        doc, "ERRATA – COORDENADORIA DE PERÍCIA MÉDICA/JOIS/COM/CBMAM",
        negrito=True, sublinhado=True, espaco_depois=8
    )

    nota_original = original.nota_bg if original else "?"
    data_pub_extenso = data_por_extenso(fechamento.data_bg_publicacao)

    _paragrafo(
        doc,
        f"ERRATA DA NOTA Nº {nota_original} – JOIS/CPM/CBMAM/{ano}, publicada "
        f"no Boletim Geral nº {fechamento.numero_bg_publicacao}, de "
        f"{data_pub_extenso}.",
        espaco_depois=10
    )

    _paragrafo(doc, "ONDE SE LÊ:", negrito=True, espaco_depois=4)
    for linha in (fechamento.onde_se_le or "").splitlines() or [""]:
        _paragrafo(doc, linha, espaco_depois=2)

    _paragrafo(doc, "", espaco_depois=6)

    _paragrafo(doc, "LEIA-SE:", negrito=True, espaco_depois=4)
    for linha in (fechamento.leia_se or "").splitlines() or [""]:
        _paragrafo(doc, linha, espaco_depois=2)

    _paragrafo(doc, "", espaco_depois=14)


# ---------------------------------------------------------------------------
# Fechamento / assinaturas
# ---------------------------------------------------------------------------

def _secao_assinaturas(doc, data_referencia: date):
    _paragrafo(
        doc, f"Perícias Médicas da JOIS/CBMAM, {data_por_extenso(data_referencia)}.",
        alinhamento=WD_ALIGN_PARAGRAPH.RIGHT, espaco_depois=24
    )

    for membro in MEMBROS_JOIS:
        _paragrafo(
            doc, f"{membro['nome']} – {membro['identificacao']}",
            negrito=False, alinhamento=WD_ALIGN_PARAGRAPH.CENTER, espaco_depois=0
        )
        _paragrafo(
            doc, membro["cargo"],
            alinhamento=WD_ALIGN_PARAGRAPH.CENTER, espaco_depois=0
        )
        _paragrafo(
            doc, membro["crm"],
            alinhamento=WD_ALIGN_PARAGRAPH.CENTER, espaco_depois=20
        )


# ---------------------------------------------------------------------------
# Montagem principal
# ---------------------------------------------------------------------------

def agrupar_por_secao(licencas: list[Licencas]):
    grupos = {
        "lts": [], "ltspf": [], "lm": [], "apto_recom": [], "apto_restr": [],
        "apto": [], "agregado": [], "curso": [], "taf": [], "promocao": [],
        "ao": [], "iso": [],
    }
    chave_por_tipo = {
        "LTS": "lts", "LTSPF": "ltspf", "LM": "lm",
        "APTO_RECOM": "apto_recom", "APTO_RESTR": "apto_restr",
        "APTO": "apto", "AGREGADO": "agregado",
        "CURSO": "curso", "TAF": "taf", "PROMOCAO": "promocao",
        "AO": "ao", "ISO": "iso",
    }
    for lic in licencas:
        chave = chave_por_tipo.get(lic.tipo_licenca)
        if chave:
            grupos[chave].append(lic)
    return grupos


def gerar_nota_bg_docx(fechamento_id: int, commit_db: bool = True) -> str:
    fechamento = (
        JuntaFechamentoBg.query
        .options(joinedload(JuntaFechamentoBg.fechamento_original))
        .get(fechamento_id)
    )

    if not fechamento:
        raise ValueError("Fechamento BG não encontrado.")

    licencas = (
        Licencas.query
        .filter_by(fechamento_bg_id=fechamento.id)
        .options(
            joinedload(Licencas.militar).joinedload(Militar.posto_grad),
            joinedload(Licencas.militar).joinedload(Militar.quadro),
            joinedload(Licencas.restricoes).joinedload(LicencaRestricao.tipo),
            joinedload(Licencas.curso),
        )
        .order_by(Licencas.tipo_licenca.asc(), Licencas.militar_id.asc())
        .all()
    )

    # Cada tabela da nota lista os BMs do mais antigo pro mais moderno na
    # hierarquia (CEL primeiro, AL SD por último) — não na ordem de
    # lançamento. `agrupar_por_secao` só distribui pra cada seção mantendo a
    # ordem recebida, e os `sorted(..., key=...)` dentro das seções de CURSO/
    # TAF/PROMOÇÃO/AGREGADO/APTO são estáveis (ordenam só por curso/data),
    # então basta ordenar a lista uma vez aqui pra ela valer em toda seção.
    licencas.sort(key=lambda lic: ordem_hierarquica_militar(lic.militar))

    grupos = agrupar_por_secao(licencas)

    doc = docx.Document(str(TEMPLATE_PATH))

    ano = fechamento.data_referencia.year

    _paragrafo(
        doc, f"NOTA PARA BG Nº {fechamento.nota_bg} JOIS/PM/CBMAM/{ano}",
        negrito=True, alinhamento=WD_ALIGN_PARAGRAPH.CENTER, espaco_depois=10
    )
    _paragrafo(
        doc, "COORDENADORIA DE PERÍCIAS MÉDICAS/JOIS/CBMAM",
        negrito=True, sublinhado=True, espaco_depois=10
    )

    if not fechamento.eh_errata:
        _paragrafo(
            doc,
            f"Relação dos bombeiros militares inspecionados pela JOIS/CBMAM, "
            f"Sessão nº{fechamento.sessao}/{ano}, no dia "
            f"{data_por_extenso(fechamento.data_referencia)}, com seus "
            f"respectivos resultados:",
            espaco_depois=12
        )

        for tipo, cabecalho in SECOES_PERIODO:
            chave = {"LTS": "lts", "LTSPF": "ltspf", "LM": "lm",
                    "APTO_RECOM": "apto_recom", "APTO_RESTR": "apto_restr"}[tipo]
            _secao_periodo(doc, tipo, cabecalho, grupos[chave])

        _secao_curso(doc, grupos["curso"])
        _secao_taf(doc, grupos["taf"])
        _secao_promocao(doc, grupos["promocao"])
        _secao_ao_iso(doc, "AO", grupos["ao"])
        _secao_ao_iso(doc, "ISO", grupos["iso"])
        _secao_agrupada_por_data(doc, grupos["agregado"], "AGREGADO")
        _secao_agrupada_por_data(doc, grupos["apto"], "APTO AO SERVIÇO DO CBMAM")
    else:
        _secao_errata(doc, fechamento)

    _secao_assinaturas(doc, fechamento.data_referencia)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    nome_arquivo = f"nota_bg_{fechamento.nota_bg.replace('/', '-')}_{fechamento.id}.docx"
    output_path = OUTPUT_DIR / nome_arquivo
    doc.save(str(output_path))

    fechamento.arquivo_docx = nome_arquivo
    if commit_db:
        database.session.commit()

    return nome_arquivo
