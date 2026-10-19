"""Consultas usadas pelas telas de listagem de situação funcional irmãs da
LTS: militares agregados, à disposição e em licença especial.

A regra de status de cada uma já existe em src.decorators.business_logic
(processar_militares_agregados/a_disposicao/le) — este módulo é só sobre
montar a listagem (com eager loading) para o template.
"""
from collections import Counter
from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import func
from sqlalchemy.orm import joinedload

from src import database
from src.models import LicencaEspecial, Militar, MilitarObmFuncao, MilitaresADisposicao, MilitaresAgregados, Obm
from src.querys import _periodo_vigente_expr
from src.services.militar_situacao_service import (
    ids_alto_comando_excluidos_de_agregado_disposicao,
    militar_com_modalidade_a_disposicao_expr,
    militar_com_situacao_agregado_expr,
)

DIAS_ALERTA_VENCIMENTO = 30
QTD_DESTINOS_NO_RANKING = 5

_RELACIONAMENTOS_COMUNS = ("militar", "posto_grad", "quadro", "destino", "modalidade", "publicacao_bg")


def _ids_mais_recentes_por_militar(model):
    """Id do registro mais novo (maior id) de cada militar_id, considerando
    só registros com `inicio_periodo` preenchido.

    Sem esse filtro, um rascunho/registro incompleto (sem início — nunca
    chegou a ser uma disposição/agregação de verdade) que por acaso tem o ID
    mais alto vira "o registro atual" do militar, mesmo que a ficha dele já
    esteja PRONTO há tempos e não tenha nada a ver com esse lixo: aconteceu
    de verdade com vários militares (quase todos com destino Defesa Civil),
    e é exatamente o tipo de dado que `encerrar_agregacao_vigente`/
    `encerrar_disposicao_vigente` fecham (`fim = ontem`) só por estarem
    "em aberto" quando a ficha muda de situação, sem nunca terem tido
    início — então continuam por aí pra sempre como o "mais recente".

    Usado pra telas de listagem mostrarem um único registro por militar,
    mesmo tendo várias linhas de histórico no banco pra ele."""
    linhas = (
        database.session.query(
            model.id,
            func.row_number().over(
                partition_by=model.militar_id,
                order_by=model.id.desc(),
            ).label("linha"),
        )
        .filter(model.inicio_periodo.isnot(None))
        .subquery()
    )
    return database.session.query(linhas.c.id).filter(linhas.c.linha == 1)


def _ids_militares_com_registro_vigente(model, campo_inicio, campo_fim, filtro_situacao_atual):
    """Ids de militar (deduplicado, sem o alto comando) com um registro
    vigente por data na tabela filha informada — mesmo critério de "Vigente"
    dos dashboards de Agregados/À Disposição. Usado pra qualquer tela que
    precise filtrar/contar "quem está à disposição (ou agregado) hoje" bater
    com o mesmo número dessas telas.

    `filtro_situacao_atual` (militar_com_modalidade_a_disposicao_expr /
    militar_com_situacao_agregado_expr) garante que o registro ainda é
    relevante AGORA — sem isso, alguém que o operador já moveu pra outro
    lugar (ex.: Agregado aguardando RR) continuaria aparecendo pra sempre só
    porque o registro antigo, já vencido, continua sendo "o mais recente"."""
    query = (
        database.session.query(model.militar_id)
        .join(Militar, Militar.id == model.militar_id)
        .filter(Militar.inativo.is_(False))
        .filter(model.id.in_(_ids_mais_recentes_por_militar(model)))
        .filter(_periodo_vigente_expr(campo_inicio, campo_fim))
        .filter(filtro_situacao_atual)
    )
    excluidos = ids_alto_comando_excluidos_de_agregado_disposicao()
    if excluidos:
        query = query.filter(model.militar_id.notin_(excluidos))
    return set(mid for (mid,) in query.all())


def ids_militares_a_disposicao_vigente():
    """Union de "só à disposição" e "agregado e à disposição ao mesmo
    tempo" — a tabela À Disposição é alimentada pela Modalidade, não pela
    Situação, então já cobre os dois casos."""
    return _ids_militares_com_registro_vigente(
        MilitaresADisposicao,
        MilitaresADisposicao.inicio_periodo,
        MilitaresADisposicao.fim_periodo_disposicao,
        militar_com_modalidade_a_disposicao_expr(),
    )


def ids_militares_agregados_vigente():
    return _ids_militares_com_registro_vigente(
        MilitaresAgregados,
        MilitaresAgregados.inicio_periodo,
        MilitaresAgregados.fim_periodo_agregacao,
        militar_com_situacao_agregado_expr(),
    )


def listar_militares_agregados(militar_id=None):
    """Sem `militar_id`, lista só o registro MAIS RECENTE de cada militar
    (tela /militares-agregados) — sem isso, um militar com histórico de
    mais de uma agregação aparecia repetido na lista, inclusive com um
    registro já vencido que a ficha dele nem mostra mais (ela só olha o
    atual), o que confundia o operador. Com `militar_id`, filtra pra um
    único militar e devolve TODO o histórico dele (módulo de Histórico)."""
    query = MilitaresAgregados.query.options(
        *[joinedload(getattr(MilitaresAgregados, rel)) for rel in _RELACIONAMENTOS_COMUNS]
    )
    if militar_id is not None:
        query = query.filter(MilitaresAgregados.militar_id == militar_id)
    else:
        query = (
            query
            .join(Militar, Militar.id == MilitaresAgregados.militar_id)
            .filter(Militar.inativo.is_(False))
            .filter(MilitaresAgregados.id.in_(_ids_mais_recentes_por_militar(MilitaresAgregados)))
            .filter(militar_com_situacao_agregado_expr())
        )
        excluidos = ids_alto_comando_excluidos_de_agregado_disposicao()
        if excluidos:
            query = query.filter(MilitaresAgregados.militar_id.notin_(excluidos))
    return query.order_by(MilitaresAgregados.fim_periodo_agregacao.desc().nullslast()).all()


def listar_militares_a_disposicao(militar_id=None):
    """Sem `militar_id`, lista só o registro MAIS RECENTE de cada militar
    (tela /militares-a-disposicao) — mesmo motivo de
    listar_militares_agregados. Com `militar_id`, filtra pra um único
    militar e devolve TODO o histórico dele (módulo de Histórico)."""
    query = MilitaresADisposicao.query.options(
        *[joinedload(getattr(MilitaresADisposicao, rel)) for rel in _RELACIONAMENTOS_COMUNS]
    )
    if militar_id is not None:
        query = query.filter(MilitaresADisposicao.militar_id == militar_id)
    else:
        query = (
            query
            .join(Militar, Militar.id == MilitaresADisposicao.militar_id)
            .filter(Militar.inativo.is_(False))
            .filter(MilitaresADisposicao.id.in_(_ids_mais_recentes_por_militar(MilitaresADisposicao)))
            .filter(militar_com_modalidade_a_disposicao_expr())
        )
        excluidos = ids_alto_comando_excluidos_de_agregado_disposicao()
        if excluidos:
            query = query.filter(MilitaresADisposicao.militar_id.notin_(excluidos))
    return query.order_by(MilitaresADisposicao.fim_periodo_disposicao.desc().nullslast()).all()


def montar_resumo_dashboard(registros, campo_fim: str, status_vencido: str = "Venceu") -> dict:
    """Cards de resumo + ranking de destinos pras telas de listagem (à
    disposição / agregados). Computado em cima da mesma lista já
    deduplicada (um registro por militar) que vai pra tabela — sem
    consulta extra ao banco.

    `campo_fim` é o nome do atributo de data de término em cada registro
    ("fim_periodo_disposicao" ou "fim_periodo_agregacao"), já que os dois
    modelos usam nomes diferentes pra essa coluna. `status_vencido` é o
    texto usado pelo status "já venceu" — MilitaresADisposicao usa "Venceu",
    mas MilitaresAgregados usa "Término de Agregação" (ver atualizar_status
    de cada model em src/models.py)."""
    hoje = date.today()
    total = len(registros)

    vigentes = a_iniciar = vencidos = vencendo_em_breve = sem_dados = 0
    contagem_destino = Counter()

    for r in registros:
        if r.status == "Vigente":
            vigentes += 1
            fim = getattr(r, campo_fim)
            if fim and 0 <= (fim - hoje).days <= DIAS_ALERTA_VENCIMENTO:
                vencendo_em_breve += 1
        elif r.status == "A iniciar":
            a_iniciar += 1
        elif r.status == status_vencido:
            vencidos += 1
        else:
            # "Inativo" (sem data de início preenchida) ou qualquer outro
            # status fora dos três normais — não pode ficar de fora da
            # soma, senão os cards não batem com o total de registros.
            sem_dados += 1

        contagem_destino[r.destino.local if r.destino else "Sem destino informado"] += 1

    top_destinos = contagem_destino.most_common(QTD_DESTINOS_NO_RANKING)
    outros = total - sum(qtd for _, qtd in top_destinos)
    if outros > 0:
        top_destinos.append(("Outros", outros))

    maior_qtd = max((qtd for _, qtd in top_destinos), default=0)
    destinos = [
        {
            "nome": nome,
            "qtd": qtd,
            "pct": round(qtd / total * 100) if total else 0,
            "largura_pct": round(qtd / maior_qtd * 100) if maior_qtd else 0,
        }
        for nome, qtd in top_destinos
    ]

    return {
        "hoje": hoje,
        "total": total,
        "vigentes": vigentes,
        "a_iniciar": a_iniciar,
        "vencidos": vencidos,
        "vencendo_em_breve": vencendo_em_breve,
        "sem_dados": sem_dados,
        "status_vencido": status_vencido,
        "destinos": destinos,
    }


def _dias_texto(status, inicio, fim, hoje, status_vencido) -> str:
    """Mesmo texto do badge "Dias" da tela (macro dias_badge em
    _situacoes_militares_macros.html), só que em texto puro pra planilha."""
    if status == "Vigente" and fim:
        dias = (fim - hoje).days
        return f"Faltam {dias} dia{'s' if dias != 1 else ''}"
    if status == status_vencido and fim:
        dias = (hoje - fim).days
        return f"Vencido há {dias} dia{'s' if dias != 1 else ''}"
    if status == "A iniciar" and inicio:
        dias = (inicio - hoje).days
        return f"Inicia em {dias} dia{'s' if dias != 1 else ''}"
    return "—"


COLUNAS_EXPORTACAO_SITUACAO = [
    "Posto/Graduação", "Matrícula", "Nome Completo", "Quadro", "Destino",
    "Situação atual do militar", "A contar de", "Término", "Dias", "Status",
    "Documento Autorizador",
]


def gerar_planilha_situacao(titulo: str, registros, campo_fim: str, status_vencido: str) -> BytesIO:
    """Gera o .xlsx de uma tela de listagem (à disposição / agregados) a
    partir da mesma lista (já deduplicada, um registro por militar) usada
    na tabela — mesmas colunas, incluindo os "Dias" calculados."""
    hoje = date.today()

    wb = Workbook()
    ws = wb.active
    ws.title = titulo[:31]  # limite de 31 caracteres do Excel pro nome da aba

    ws.append(COLUNAS_EXPORTACAO_SITUACAO)
    for col_num in range(1, len(COLUNAS_EXPORTACAO_SITUACAO) + 1):
        ws.cell(row=1, column=col_num).font = Font(bold=True)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUNAS_EXPORTACAO_SITUACAO))}1"

    def fmt(dt):
        return dt.strftime("%d/%m/%Y") if dt else ""

    for r in registros:
        fim = getattr(r, campo_fim)
        ws.append([
            r.posto_grad.sigla if r.posto_grad else "",
            r.militar.matricula if r.militar else "",
            r.militar.nome_completo if r.militar else "",
            r.quadro.quadro if r.quadro else "",
            r.destino.local if r.destino else "",
            r.militar.situacao if r.militar and r.militar.situacao else "",
            fmt(r.inicio_periodo),
            fmt(fim),
            _dias_texto(r.status, r.inicio_periodo, fim, hoje, status_vencido),
            r.status or "",
            r.publicacao_bg.boletim_geral if r.publicacao_bg else "",
        ])

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


# Colunas que podem ir no Excel de /militares-a-disposicao, agrupadas como no
# modal "Personalizar planilha" (mesma ideia do Mapa da Força). A ordem aqui é
# a ordem das colunas no arquivo.
GRUPOS_COLUNAS_DISPOSICAO = [
    ("Identificação", [
        "Posto/Graduação", "Nome Completo", "Nome de Guerra", "Matrícula", "CPF", "RG",
        "Quadro", "Sexo", "Data de Nascimento", "Inclusão", "Especialidade", "Localidade",
    ]),
    ("Lotação", ["OBM atual", "Função atual"]),
    ("Disposição", [
        "Destino", "Modalidade", "Situação atual do militar", "A contar de", "Término",
        "Dias", "Status", "Documento Autorizador",
    ]),
    ("Situação no momento", ["Também em situação extra", "De férias agora"]),
]
COLUNAS_DISPOSICAO_DISPONIVEIS = [c for _, cols in GRUPOS_COLUNAS_DISPOSICAO for c in cols]
# marcadas por padrão no modal (as que a tabela da tela mostra + OBM e badges)
COLUNAS_DISPOSICAO_PADRAO = [
    "Posto/Graduação", "Nome Completo", "Matrícula", "Quadro", "OBM atual", "Destino",
    "Situação atual do militar", "A contar de", "Término", "Dias", "Status",
    "Documento Autorizador", "Também em situação extra", "De férias agora",
]
ABAS_EXTRAS_DISPOSICAO = ["Resumo", "Por destino", "Pendências"]
STATUS_DISPOSICAO = ["Vigente", "A iniciar", "Venceu", "Inativo"]


def _lotacao_atual_por_militar(militar_ids) -> dict:
    """OBM(s) e função(ões) com vínculo ativo de cada militar, numa consulta só."""
    if not militar_ids:
        return {}
    from src.models import Funcao
    linhas = (
        database.session.query(MilitarObmFuncao.militar_id, Obm.sigla, Funcao.ocupacao)
        .outerjoin(Obm, Obm.id == MilitarObmFuncao.obm_id)
        .outerjoin(Funcao, Funcao.id == MilitarObmFuncao.funcao_id)
        .filter(MilitarObmFuncao.militar_id.in_(set(militar_ids)))
        .filter(MilitarObmFuncao.data_fim.is_(None))
        .all()
    )
    lotacao = {}
    for mid, sigla, funcao in linhas:
        item = lotacao.setdefault(mid, {"obms": [], "funcoes": []})
        if sigla and sigla not in item["obms"]:
            item["obms"].append(sigla)
        if funcao and funcao not in item["funcoes"]:
            item["funcoes"].append(funcao)
    return {mid: {"obm": " / ".join(v["obms"]), "funcao": " / ".join(v["funcoes"])}
            for mid, v in lotacao.items()}


def _aba(wb, titulo, colunas, linhas, primeira=False):
    """Cria uma aba com cabeçalho em destaque, filtro, painel congelado e
    largura de coluna ajustada ao conteúdo."""
    ws = wb.active if primeira else wb.create_sheet()
    ws.title = titulo[:31]
    ws.append(colunas)
    for linha in linhas:
        ws.append(linha)
    destaque = PatternFill("solid", fgColor="0B385E")
    for cel in ws[1]:
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = destaque
        cel.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(colunas))}{max(len(linhas) + 1, 1)}"
    for i, col in enumerate(colunas, start=1):
        maior = max([len(str(col))] + [len(str(l[i - 1])) for l in linhas[:500] if l[i - 1] is not None])
        ws.column_dimensions[get_column_letter(i)].width = min(maior + 2, 50)
    return ws


def gerar_planilha_disposicao_completa(registros, pendencias, situacoes_extras_por_militar,
                                       ferias_por_militar, colunas=None, abas_extras=None) -> BytesIO:
    """Excel da tela /militares-a-disposicao montado pelo usuário: só as
    `colunas` escolhidas (na ordem de GRUPOS_COLUNAS_DISPOSICAO) e, se pedidas,
    as abas extras de resumo, contagem por destino e pendências."""
    colunas = [c for c in COLUNAS_DISPOSICAO_DISPONIVEIS if c in (colunas or COLUNAS_DISPOSICAO_PADRAO)]
    abas_extras = ABAS_EXTRAS_DISPOSICAO if abas_extras is None else abas_extras
    hoje = date.today()
    fmt = lambda d: d.strftime("%d/%m/%Y") if d else ""
    lotacao = _lotacao_atual_por_militar([r.militar_id for r in registros])

    def sexo(m):
        s = (m.sexo or "").strip().lower()
        return "Masculino" if s.startswith("m") else "Feminino" if s.startswith("f") else (m.sexo or "")

    linhas = []
    for r in registros:
        m = r.militar
        extra = situacoes_extras_por_militar.get(r.militar_id)
        ferias = ferias_por_militar.get(r.militar_id)
        texto_extra = ""
        if extra:
            texto_extra = extra["label"]
            if extra.get("destino"):
                texto_extra += f" — {extra['destino']}"
            texto_extra += f" ({fmt(extra['inicio'])} a {fmt(extra['fim']) or 'sem término'})"
        lot = lotacao.get(r.militar_id, {})
        valores = {
            "Posto/Graduação": r.posto_grad.sigla if r.posto_grad else "",
            "Nome Completo": m.nome_completo if m else "",
            "Nome de Guerra": (m.nome_guerra or "") if m else "",
            "Matrícula": (m.matricula or "") if m else "",
            "CPF": (m.cpf or "") if m else "",
            "RG": (m.rg or "") if m else "",
            "Quadro": r.quadro.quadro if r.quadro else "",
            "Sexo": sexo(m) if m else "",
            "Data de Nascimento": fmt(m.data_nascimento) if m else "",
            "Inclusão": fmt(m.inclusao) if m else "",
            "Especialidade": m.especialidade.ocupacao if m and m.especialidade else "",
            "Localidade": m.localidade.sigla if m and m.localidade else "",
            "OBM atual": lot.get("obm", ""),
            "Função atual": lot.get("funcao", ""),
            "Destino": r.destino.local if r.destino else "",
            "Modalidade": r.modalidade.descricao if r.modalidade else
                          (m.modalidade.descricao if m and m.modalidade else ""),
            "Situação atual do militar": (m.situacao or "") if m else "",
            "A contar de": fmt(r.inicio_periodo),
            "Término": fmt(r.fim_periodo_disposicao),
            "Dias": _dias_texto(r.status, r.inicio_periodo, r.fim_periodo_disposicao, hoje, "Venceu"),
            "Status": r.status or "",
            "Documento Autorizador": r.publicacao_bg.boletim_geral if r.publicacao_bg else "",
            "Também em situação extra": texto_extra,
            "De férias agora": (f"{ferias['numero']}º período: {fmt(ferias['inicio'])} a {fmt(ferias['fim'])}"
                                if ferias else ""),
        }
        linhas.append([valores[c] for c in colunas])

    wb = Workbook()
    _aba(wb, "Militares à disposição", colunas, linhas, primeira=True)

    if "Resumo" in abas_extras:
        resumo = montar_resumo_dashboard(registros, campo_fim="fim_periodo_disposicao")
        _aba(wb, "Resumo", ["Indicador", "Quantidade"], [
            ["Data da extração", hoje.strftime("%d/%m/%Y")],
            ["Total de militares na planilha", resumo["total"]],
            ["Vigentes", resumo["vigentes"]],
            [f"Vencendo em até {DIAS_ALERTA_VENCIMENTO} dias", resumo["vencendo_em_breve"]],
            ["A iniciar", resumo["a_iniciar"]],
            ["Vencidos", resumo["vencidos"]],
            ["Sem dados (sem data de início)", resumo["sem_dados"]],
            ["Pendências (vencida aguardando decisão)", len(pendencias)],
        ])

    if "Por destino" in abas_extras:
        total = len(registros)
        por_destino = Counter(r.destino.local if r.destino else "Sem destino informado" for r in registros)
        _aba(wb, "Por destino", ["Destino", "Militares", "% do total"],
             [[nome, qtd, f"{round(qtd / total * 100) if total else 0}%"] for nome, qtd in por_destino.most_common()])

    if "Pendências" in abas_extras:
        linhas_pend = []
        for p in pendencias:
            m, d = p["militar"], p["disposicao"]
            linhas_pend.append([
                d.posto_grad.sigla if d.posto_grad else "",
                m.matricula or "",
                m.nome_completo,
                d.destino.local if d.destino else "",
                fmt(d.fim_periodo_disposicao),
                p["dias_vencido"],
                "Sim" if p["dual"] else "Não",
                d.publicacao_bg.boletim_geral if d.publicacao_bg else "",
            ])
        _aba(wb, "Pendências", ["Posto/Graduação", "Matrícula", "Nome Completo", "Destino", "Término",
                                "Dias vencido", "Tem agregação aberta", "Documento Autorizador"], linhas_pend)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def listar_licencas_especiais(militar_id=None):
    """Sem `militar_id`, lista todas (tela /licenca-especial). Com
    `militar_id`, filtra pra um único militar (módulo de Histórico)."""
    query = LicencaEspecial.query.options(
        *[joinedload(getattr(LicencaEspecial, rel)) for rel in _RELACIONAMENTOS_COMUNS]
    )
    if militar_id is not None:
        query = query.filter(LicencaEspecial.militar_id == militar_id)
    return query.order_by(LicencaEspecial.fim_periodo_le.desc().nullslast()).all()
