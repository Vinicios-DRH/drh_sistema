from io import BytesIO
from math import ceil
from pathlib import Path

from flask import Blueprint, jsonify, render_template, request, redirect, send_file, url_for, flash, send_from_directory, session as flask_session
from flask_login import login_required, current_user
from sqlalchemy import or_, case
from sqlalchemy.orm import joinedload

from src import database
from src.authz import require_perm
from src.models import (
    Curso,
    JuntaFechamentoBg,
    JuntaRestricaoTipo,
    LicencaRestricao,
    Licencas,
    Militar,
    Obm,
    PostoGrad,
    Quadro,
)
from src.forms import FormLicencas
from src.services.junta_estatisticas import (
    montar_estatisticas_mensais,
    montar_ranking_militares,
    normalizar_competencia,
)
from src.services.junta_medica import (
    calcular_data_fim,
    calcular_situacao_atual,
    calcular_status_registro,
    contar_efetivo_ativo,
    exige_inspecao_pos_lts,
    label_status,
    label_tipo,
    listar_militares_aptos_pendentes,
    listar_sessoes_pendentes,
    listar_tipos_restricao,
    montar_dados_licencas,
    obter_ou_criar_tipo_restricao,
    ordem_hierarquica_militar,
    resultado_valido,
    RESULTADOS_COM_DETALHE_LIVRE,
    RESULTADOS_POR_TIPO,
    STATUS_LABELS,
    TIPO_LICENCA_LABELS,
    TIPOS_COM_RESTRICAO,
    TIPOS_DATA_UNICA,
    TIPOS_PONTUAIS,
)
from src.services.junta_bg_generator import gerar_nota_bg_docx
from src.services.junta_periodos import montar_blocos_por_militar
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime, date, timedelta
from calendar import monthrange
from zoneinfo import ZoneInfo

junta_bp = Blueprint("junta", __name__, url_prefix="/junta")

MANAUS_TZ = ZoneInfo("America/Manaus")


def hoje_manaus():
    return datetime.now(MANAUS_TZ).date()


def _obter_obm_atual(militar):
    obm_sigla = ""

    if hasattr(militar, "obm_funcoes") and militar.obm_funcoes:
        obms_ativas = [x for x in militar.obm_funcoes if getattr(
            x, "data_fim", None) is None]

        if obms_ativas:
            obm = obms_ativas[-1].obm
            obm_sigla = obm.sigla if obm else ""
        else:
            obm = militar.obm_funcoes[-1].obm
            obm_sigla = obm.sigla if obm else ""

    return obm_sigla or ""


MESES_PT = [
    "JANEIRO", "FEVEREIRO", "MARÇO", "ABRIL", "MAIO", "JUNHO",
    "JULHO", "AGOSTO", "SETEMBRO", "OUTUBRO", "NOVEMBRO", "DEZEMBRO"
]


def data_por_extenso_maiuscula(dt):
    if not dt:
        return ""
    return f"{dt.day} DE {MESES_PT[dt.month - 1]} DE {dt.year}"


def _listar_cursos():
    """Catálogo de cursos pra montar o select de 'Qual o curso?'."""
    return Curso.query.order_by(Curso.nome.asc()).all()


def _ids_restricoes_do_form():
    """Ids marcados nos checkboxes de tipo de restrição."""
    ids = []
    for valor in request.form.getlist("restricoes"):
        try:
            ids.append(int(valor))
        except (TypeError, ValueError):
            continue
    return ids


def _aplicar_restricoes(licenca, ids_tipos, restricao_outra=""):
    """
    Vincula as restrições marcadas ao lançamento, sem duplicar as que já
    estiverem lá. Devolve quantas foram efetivamente adicionadas.
    """
    ja_vinculadas = {r.restricao_tipo_id for r in licenca.restricoes}
    adicionadas = 0

    tipos = list(ids_tipos)

    novo_tipo = obter_ou_criar_tipo_restricao(restricao_outra)
    if novo_tipo is not None:
        tipos.append(novo_tipo.id)

    for tipo_id in tipos:
        if tipo_id in ja_vinculadas:
            continue

        database.session.add(LicencaRestricao(
            licenca_id=licenca.id,
            restricao_tipo_id=tipo_id,
        ))
        ja_vinculadas.add(tipo_id)
        adicionadas += 1

    return adicionadas


# Chaves do cookie de sessão do Flask (current_user já usa "session" pra
# outra coisa — daqui pra baixo, "sessão ativa" é sempre isto: a sessão da
# JUNTA que o operador está preenchendo agora, não a sessão de login).
_CHAVE_SESSAO_ATIVA = "junta_sessao_ativa"
_CHAVE_DATA_SESSAO_ATIVA = "junta_data_sessao_ativa"


def _sessao_ativa():
    """
    Sessão da Junta "em andamento" pra este operador (guardada no cookie).
    Preenchida uma vez no primeiro lançamento e reaproveitada nos seguintes,
    até ele trocar — é o que evita digitar sessão/data pra cada militar.
    """
    sessao = flask_session.get(_CHAVE_SESSAO_ATIVA)
    data_str = flask_session.get(_CHAVE_DATA_SESSAO_ATIVA)

    if not sessao or not data_str:
        return None

    try:
        data_sessao = datetime.strptime(data_str, "%Y-%m-%d").date()
    except ValueError:
        return None

    return {"sessao": sessao, "data_sessao": data_sessao}


def _definir_sessao_ativa(sessao: str, data_sessao):
    flask_session[_CHAVE_SESSAO_ATIVA] = sessao
    flask_session[_CHAVE_DATA_SESSAO_ATIVA] = data_sessao.isoformat()


def _limpar_sessao_ativa():
    flask_session.pop(_CHAVE_SESSAO_ATIVA, None)
    flask_session.pop(_CHAVE_DATA_SESSAO_ATIVA, None)


def _contexto_nova_licenca(form, hoje, data_extenso_hoje, sessao_ativa=None,
                           restricoes_marcadas=None):
    restricoes_marcadas = restricoes_marcadas or set()
    sessao_ja_fechada = False
    if sessao_ativa:
        # Aviso pra quem esquece de trocar de sessão: se essa mesma dupla
        # sessão+data já tem algum lançamento FECHADO, um lançamento novo
        # aqui vai abrir uma nota adicional pra ela — não vai entrar na nota
        # já gerada. Se for correção, o caminho é a errata.
        sessao_ja_fechada = (
            Licencas.query
            .filter_by(
                sessao=sessao_ativa["sessao"],
                data_sessao=sessao_ativa["data_sessao"]
            )
            .filter(Licencas.fechamento_bg_id.isnot(None))
            .first() is not None
        )

    return dict(
        form=form,
        tipo_labels=TIPO_LICENCA_LABELS,
        status_labels=STATUS_LABELS,
        sessoes_pendentes=listar_sessoes_pendentes(),
        sessao_ativa=sessao_ativa,
        sessao_ja_fechada=sessao_ja_fechada,
        hoje=hoje,
        data_extenso_hoje=data_extenso_hoje,
        cursos=_listar_cursos(),
        tipos_restricao=listar_tipos_restricao(),
        resultados_por_tipo=RESULTADOS_POR_TIPO,
        resultados_com_detalhe_livre=sorted(RESULTADOS_COM_DETALHE_LIVRE),
        tipos_pontuais=sorted(TIPOS_PONTUAIS),
        tipos_com_restricao=sorted(TIPOS_COM_RESTRICAO),
        tipos_com_reavaliar=sorted(("LTS", "APTO_RESTR", "APTO_RECOM")),
        tipos_data_unica=sorted(TIPOS_DATA_UNICA),
        aptos_pendentes=_listar_aptos_pendentes_para_tela(),
        fechamentos_anteriores=_listar_fechamentos_para_errata(),
        lancamentos_pendentes=_listar_lancamentos_pendentes(),
        restricoes_marcadas=restricoes_marcadas,
    )


def _listar_lancamentos_pendentes():
    """
    Todo lançamento ainda sem nota de BG, com militar/tipo/restrições — é a
    conferência antes de fechar qualquer sessão. Cada linha tem ação de
    editar ou excluir, pra corrigir engano (ex.: LTS lançada pro militar
    errado) antes de virar nota — depois disso o histórico é imutável.
    """
    registros = (
        Licencas.query
        .filter(Licencas.fechamento_bg_id.is_(None))
        .options(
            joinedload(Licencas.militar).joinedload(Militar.posto_grad),
            joinedload(Licencas.militar).joinedload(Militar.quadro),
            joinedload(Licencas.restricoes).joinedload(LicencaRestricao.tipo),
        )
        .order_by(Licencas.data_sessao.desc(), Licencas.sessao.desc(),
                 Licencas.created_at.asc())
        .all()
    )

    itens = []
    for reg in registros:
        itens.append({
            "registro": reg,
            "tipo_label": label_tipo(reg.tipo_licenca),
            "restricoes": sorted(
                (v.tipo.nome for v in reg.restricoes if v.tipo), key=str.lower),
        })
    return itens


def _listar_aptos_pendentes_para_tela():
    """
    Mesma lista de `listar_militares_aptos_pendentes`, mas já com a data em
    que o militar passou a ser considerado apto — Jinja não faz conta de
    data, então isso vem pronto do Python.
    """
    itens = []
    for reg in listar_militares_aptos_pendentes():
        itens.append({
            "registro": reg,
            "militar": reg.militar,
            "apto_desde": reg.data_fim + timedelta(days=1),
        })
    return itens


def _listar_fechamentos_para_errata(limite: int = 60):
    """Notas de BG já finalizadas — pra escolher qual será corrigida na errata."""
    return (
        JuntaFechamentoBg.query
        .filter(JuntaFechamentoBg.eh_errata.is_(False))
        .order_by(JuntaFechamentoBg.data_referencia.desc(),
                 JuntaFechamentoBg.id.desc())
        .limit(limite)
        .all()
    )


class _ErroValidacaoLicenca(Exception):
    """Erro de validação específico do tipo de inspeção (flash pronto)."""


def _calcular_campos_por_tipo(form, tipo, data_sessao, status_atual):
    """
    Calcula os campos que dependem do tipo de inspeção — dias, datas, status,
    curso, resultado, reavaliação. Usado tanto ao CRIAR quanto ao EDITAR um
    lançamento, pra não duplicar essa lógica (e correr o risco de um dia
    corrigir só num lugar).
    """
    campos = {
        "numero_bg_curso": None,
        "data_extenso_curso": None,
        "curso_id": None,
        "curso_nome": None,
        "resultado_detalhe": None,
        # Só é lido pra LTS/APTO_RESTR/APTO_RECOM — pros demais tipos esse
        # valor nunca é consultado.
        "reavaliar_ao_termino": True,
        "online": bool(form.online.data),
    }

    if tipo in TIPOS_PONTUAIS:
        # CURSO / TAF / PROMOÇÃO: parecer pontual, com resultado próprio e
        # valendo pela data da sessão da Junta.
        resultado = (form.resultado_inspecao.data or "").strip()

        if not resultado_valido(tipo, resultado):
            raise _ErroValidacaoLicenca(
                f"Informe um resultado válido para a inspeção de {label_tipo(tipo)}.")

        if resultado in RESULTADOS_COM_DETALHE_LIVRE:
            resultado_detalhe = (form.resultado_detalhe.data or "").strip()

            if not resultado_detalhe:
                raise _ErroValidacaoLicenca(
                    "Informe qual foi o resultado (opção 'Outro').")

            campos["resultado_detalhe"] = resultado_detalhe

        if tipo == "CURSO":
            curso_id, curso_nome = _resolver_curso(form)

            if not curso_nome:
                raise _ErroValidacaoLicenca("Informe qual o curso da inspeção.")

            numero_bg_curso = (form.numero_bg_curso.data or "").strip()

            if not numero_bg_curso:
                raise _ErroValidacaoLicenca(
                    "Informe o número do BG para fins de curso.")

            campos["curso_id"] = curso_id
            campos["curso_nome"] = curso_nome
            campos["numero_bg_curso"] = numero_bg_curso
        else:
            campos["numero_bg_curso"] = (
                form.numero_bg_curso.data or "").strip() or None

        campos["qtd_dias"] = 1
        campos["data_inicio"] = data_sessao
        campos["data_fim"] = data_sessao
        campos["status_registro"] = resultado
        campos["data_extenso_curso"] = data_por_extenso_maiuscula(data_sessao)

    elif tipo == "AGREGADO":
        data_inicio = form.data_inicio.data or data_sessao

        if status_atual != "APTO_RESTR":
            raise _ErroValidacaoLicenca(
                "A agregação manual só pode ser registrada quando a situação "
                "atual do militar estiver como APTO COM RESTRIÇÕES.")

        campos["qtd_dias"] = 1
        campos["data_inicio"] = data_inicio
        campos["data_fim"] = data_inicio
        campos["status_registro"] = calcular_status_registro(tipo)

    elif tipo == "APTO":
        # Só uma data: o militar está apto a partir dela, sem prazo.
        data_inicio = form.data_inicio.data

        if not data_inicio:
            raise _ErroValidacaoLicenca("Informe a data.")

        campos["qtd_dias"] = 1
        campos["data_inicio"] = data_inicio
        campos["data_fim"] = data_inicio
        campos["status_registro"] = calcular_status_registro(tipo)

    else:
        data_inicio = form.data_inicio.data
        qtd_dias = form.qtd_dias.data

        if not data_inicio:
            raise _ErroValidacaoLicenca("Informe a data de início.")

        if not qtd_dias:
            raise _ErroValidacaoLicenca("Informe a quantidade de dias.")

        campos["qtd_dias"] = qtd_dias
        campos["data_inicio"] = data_inicio
        campos["data_fim"] = calcular_data_fim(data_inicio, qtd_dias)
        campos["status_registro"] = calcular_status_registro(tipo)

        # O checkbox também vale pra APTO_RESTR/APTO_RECOM: é o mesmo padrão
        # que a nota oficial da JOIS usa nesses pareceres ("REAVALIAR AO
        # TÉRMINO" x "PRONTO PARA SV"), então a tela deixa marcar nos três.
        # Pra LTSPF/LM não existe essa decisão — passado o prazo, sempre
        # voltam a ser aptos sozinhos.
        if tipo in ("LTS", "APTO_RESTR", "APTO_RECOM"):
            campos["reavaliar_ao_termino"] = bool(
                form.reavaliar_ao_termino.data)

    return campos


@junta_bp.route("/nova-licenca", methods=["GET", "POST"])
@login_required
@require_perm("JUNTA_CREATE")
def nova_licenca():
    if request.method == "GET" and request.args.get("trocar_sessao"):
        _limpar_sessao_ativa()
        return redirect(url_for("junta.nova_licenca"))

    form = FormLicencas()
    hoje = hoje_manaus()
    data_extenso_hoje = data_por_extenso_maiuscula(hoje)

    sessao_ativa = _sessao_ativa()

    if request.method == "GET" and sessao_ativa:
        form.sessao.data = sessao_ativa["sessao"]
        form.data_sessao.data = sessao_ativa["data_sessao"]

    if form.validate_on_submit():
        try:
            militar_id = int(form.militar_id.data)
            militar = Militar.query.get(militar_id)

            if not militar:
                flash("Militar não encontrado.", "danger")
                return redirect(url_for("junta.nova_licenca"))

            data_sessao = form.data_sessao.data
            tipo = form.tipo_licenca.data

            historico = (
                Licencas.query
                .filter_by(militar_id=militar.id)
                .order_by(Licencas.data_inicio.desc(), Licencas.id.desc())
                .all()
            )

            situacao = calcular_situacao_atual(historico)
            status_atual = situacao["status_atual"]

            campos = _calcular_campos_por_tipo(
                form, tipo, data_sessao, status_atual)

            nova = Licencas(
                militar_id=militar.id,
                tipo_licenca=tipo,
                recebimento_bg=None,
                qtd_dias=campos["qtd_dias"],
                data_inicio=campos["data_inicio"],
                data_fim=campos["data_fim"],
                status=campos["status_registro"],
                sessao=form.sessao.data.strip(),
                data_sessao=data_sessao,
                numero_bg_curso=campos["numero_bg_curso"],
                data_extenso_curso=campos["data_extenso_curso"],
                curso_id=campos["curso_id"],
                curso_nome=campos["curso_nome"],
                resultado_detalhe=campos["resultado_detalhe"],
                reavaliar_ao_termino=campos["reavaliar_ao_termino"],
                online=campos["online"],
                observacao=form.observacao.data.strip() if form.observacao.data else None,
                usuario_id=current_user.id
            )

            database.session.add(nova)
            database.session.flush()

            if tipo in TIPOS_COM_RESTRICAO:
                _aplicar_restricoes(
                    nova,
                    _ids_restricoes_do_form(),
                    form.restricao_outra.data or ""
                )

            database.session.commit()

            # Trava a sessão pro próximo lançamento — só precisa digitar de
            # novo se clicar em "Trocar sessão".
            _definir_sessao_ativa(form.sessao.data.strip(), data_sessao)

            flash("Registro da Junta Médica adicionado com sucesso!", "success")
            return redirect(url_for("junta.nova_licenca"))

        except _ErroValidacaoLicenca as e:
            database.session.rollback()
            flash(str(e), "danger")
        except Exception as e:
            database.session.rollback()
            flash(f"Erro ao salvar licença: {str(e)}", "danger")

    elif request.method == "POST":
        for campo, erros in form.errors.items():
            rotulo = getattr(form, campo).label.text if hasattr(
                form, campo) else campo
            flash(f"{rotulo}: {'; '.join(erros)}", "danger")

    return render_template(
        "junta/nova_licenca.html",
        **_contexto_nova_licenca(form, hoje, data_extenso_hoje,
                                 sessao_ativa=sessao_ativa)
    )


@junta_bp.route("/licenca/<int:licenca_id>/editar", methods=["GET", "POST"])
@login_required
@require_perm("JUNTA_CREATE")
def editar_licenca(licenca_id):
    """
    Corrige um lançamento pendente — ex.: LTS lançada pro militar errado.
    Só é permitido enquanto o lançamento não tiver entrado numa nota de BG
    (histórico vira imutável a partir daí; depois disso é errata).
    """
    licenca = Licencas.query.get_or_404(licenca_id)

    if licenca.fechamento_bg_id is not None:
        flash(
            "Este lançamento já foi fechado numa nota de BG e não pode mais "
            "ser editado. Se for pra corrigir algo já publicado, use errata.",
            "warning"
        )
        return redirect(url_for("junta.listar_licencas"))

    form = FormLicencas(obj=licenca)
    hoje = hoje_manaus()
    data_extenso_hoje = data_por_extenso_maiuscula(hoje)

    if request.method == "GET":
        # `obj=licenca` já populou o que tem nome igual à coluna (sessao,
        # data_sessao, tipo_licenca, qtd_dias, data_inicio, observacao,
        # numero_bg_curso, resultado_detalhe, reavaliar_ao_termino, online).
        # O resto precisa de ajuste manual porque não bate 1:1 com o modelo.
        form.militar_id.data = str(licenca.militar_id)
        form.militar_nome.data = licenca.militar.nome_completo
        form.posto_grad_id.data = (
            licenca.militar.posto_grad.sigla if licenca.militar.posto_grad else "")
        form.quadro_id.data = (
            licenca.militar.quadro.quadro if licenca.militar.quadro else "")
        form.obm_id_1.data = _obter_obm_atual(licenca.militar)

        if licenca.tipo_licenca in TIPOS_PONTUAIS:
            form.resultado_inspecao.data = licenca.status

        if licenca.tipo_licenca == "CURSO":
            if licenca.curso_id:
                form.curso_id.data = str(licenca.curso_id)
            elif licenca.curso_nome:
                form.curso_id.data = "OUTRO"
                form.curso_outro.data = licenca.curso_nome

    if form.validate_on_submit():
        try:
            militar_id = int(form.militar_id.data)
            militar = Militar.query.get(militar_id)

            if not militar:
                flash("Militar não encontrado.", "danger")
                return redirect(url_for("junta.editar_licenca", licenca_id=licenca.id))

            data_sessao = form.data_sessao.data
            tipo = form.tipo_licenca.data

            historico = (
                Licencas.query
                .filter(Licencas.militar_id == militar.id,
                       Licencas.id != licenca.id)
                .order_by(Licencas.data_inicio.desc(), Licencas.id.desc())
                .all()
            )

            situacao = calcular_situacao_atual(historico)
            status_atual = situacao["status_atual"]

            campos = _calcular_campos_por_tipo(
                form, tipo, data_sessao, status_atual)

            licenca.militar_id = militar.id
            licenca.tipo_licenca = tipo
            licenca.qtd_dias = campos["qtd_dias"]
            licenca.data_inicio = campos["data_inicio"]
            licenca.data_fim = campos["data_fim"]
            licenca.status = campos["status_registro"]
            licenca.sessao = form.sessao.data.strip()
            licenca.data_sessao = data_sessao
            licenca.numero_bg_curso = campos["numero_bg_curso"]
            licenca.data_extenso_curso = campos["data_extenso_curso"]
            licenca.curso_id = campos["curso_id"]
            licenca.curso_nome = campos["curso_nome"]
            licenca.resultado_detalhe = campos["resultado_detalhe"]
            licenca.reavaliar_ao_termino = campos["reavaliar_ao_termino"]
            licenca.online = campos["online"]
            licenca.observacao = form.observacao.data.strip() if form.observacao.data else None

            # Restrições: limpa os vínculos antigos e reaplica do zero — mais
            # simples e seguro do que tentar diferenciar checkbox por
            # checkbox. `.clear()` na coleção (não uma query bulk) é de
            # propósito: o relacionamento tem cascade="delete-orphan", então
            # o SQLAlchemy apaga as linhas órfãs sozinho E mantém o estado em
            # memória sincronizado — importante porque `_aplicar_restricoes`
            # olha `licenca.restricoes` em seguida pra não duplicar.
            licenca.restricoes.clear()
            database.session.flush()

            if tipo in TIPOS_COM_RESTRICAO:
                _aplicar_restricoes(
                    licenca,
                    _ids_restricoes_do_form(),
                    form.restricao_outra.data or ""
                )

            database.session.commit()

            flash("Lançamento atualizado com sucesso!", "success")
            return redirect(url_for("junta.nova_licenca"))

        except _ErroValidacaoLicenca as e:
            database.session.rollback()
            flash(str(e), "danger")
        except Exception as e:
            database.session.rollback()
            flash(f"Erro ao atualizar lançamento: {str(e)}", "danger")

    elif request.method == "POST":
        for campo, erros in form.errors.items():
            rotulo = getattr(form, campo).label.text if hasattr(
                form, campo) else campo
            flash(f"{rotulo}: {'; '.join(erros)}", "danger")

    restricoes_marcadas = {v.restricao_tipo_id for v in licenca.restricoes}

    return render_template(
        "junta/nova_licenca.html",
        **_contexto_nova_licenca(form, hoje, data_extenso_hoje, sessao_ativa=None,
                                 restricoes_marcadas=restricoes_marcadas),
        modo_edicao=True,
        licenca_id=licenca.id,
    )


@junta_bp.route("/licenca/<int:licenca_id>/excluir", methods=["POST"])
@login_required
@require_perm("JUNTA_CREATE")
def excluir_licenca(licenca_id):
    """Remove um lançamento pendente por engano — só antes de virar nota."""
    licenca = Licencas.query.get_or_404(licenca_id)

    if licenca.fechamento_bg_id is not None:
        flash(
            "Este lançamento já foi fechado numa nota de BG e não pode mais "
            "ser excluído.", "warning"
        )
        return redirect(url_for("junta.nova_licenca"))

    try:
        militar_nome = licenca.militar.nome_completo
        # A relação tem cascade="delete-orphan" — apagar a licença já apaga
        # as restrições vinculadas a ela sozinho.
        database.session.delete(licenca)
        database.session.commit()
        flash(f"Lançamento de {militar_nome} excluído.", "success")
    except Exception as e:
        database.session.rollback()
        flash(f"Erro ao excluir lançamento: {str(e)}", "danger")

    return redirect(url_for("junta.nova_licenca"))


def _resolver_curso(form):
    """
    Resolve o curso da inspeção: catálogo (`curso_id`) ou o nome digitado em
    "Outros". O nome sempre volta preenchido — é ele que vai pra nota do BG.
    """
    curso_outro = (form.curso_outro.data or "").strip()
    escolha = (form.curso_id.data or "").strip()

    if escolha and escolha != "OUTRO":
        try:
            curso = Curso.query.get(int(escolha))
        except (TypeError, ValueError):
            curso = None

        if curso:
            return curso.id, curso.nome

    return None, (curso_outro or None)


@junta_bp.route("/licenca/<int:licenca_id>/restricoes", methods=["POST"])
@login_required
@require_perm("JUNTA_CREATE")
def adicionar_restricoes(licenca_id):
    """
    Acrescenta restrições a um lançamento que já existe — é como o militar
    ganha restrições novas sem precisar refazer o parecer inteiro.
    """
    licenca = Licencas.query.get_or_404(licenca_id)
    origem = request.form.get("origem") or url_for("junta.listar_licencas")

    try:
        adicionadas = _aplicar_restricoes(
            licenca,
            _ids_restricoes_do_form(),
            request.form.get("restricao_outra") or ""
        )
        database.session.commit()

        if adicionadas:
            flash(
                f"{adicionadas} restrição(ões) adicionada(s) ao militar.", "success")
        else:
            flash("Nenhuma restrição nova para adicionar.", "warning")

    except Exception as e:
        database.session.rollback()
        flash(f"Erro ao adicionar restrições: {str(e)}", "danger")

    return redirect(origem)


@junta_bp.route("/licencas", methods=["GET"])
@login_required
@require_perm("JUNTA_READ")
def listar_licencas():
    page = request.args.get("page", 1, type=int)
    per_page = 20

    filtros = _ler_filtros_listagem()

    dados, resumo = montar_dados_licencas(**filtros)

    total = len(dados)
    total_pages = max(1, ceil(total / per_page)) if total else 1

    if page > total_pages:
        page = total_pages

    inicio = (page - 1) * per_page
    fim = inicio + per_page
    dados_pagina = dados[inicio:fim]

    tipos_filtro = [
        ("LTS", "LTS"),
        ("LTSPF", "LTSPF"),
        ("LM", "Licença Maternidade"),
        ("APTO_RECOM", "Apto com Recomendações"),
        ("APTO_RESTR", "Apto com Restrições"),
        ("APTO", "Apto sem Restrição"),
        ("CURSO", "Curso"),
        ("TAF", "TAF"),
        ("PROMOCAO", "Promoção"),
        ("AGREGADO", "Agregado"),
    ]

    status_filtro = [
        ("LTS", "LTS"),
        ("LTSPF", "LTSPF"),
        ("LM", "Licença Maternidade"),
        ("APTO_RECOM", "Apto com Recomendações"),
        ("APTO_RESTR", "Apto com Restrições"),
        ("APTO", "Apto sem Restrição"),
        ("CURSO_APTO", "Curso — Apto"),
        ("CURSO_REGIME_ESPECIAL", "Curso — Regime Especial"),
        ("CURSO_INAPTO", "Curso — Inapto"),
        ("TAF_APTO", "TAF — Apto"),
        ("TAF_ALTERNATIVO", "TAF — Alternativo"),
        ("TAF_INAPTO", "TAF — Inapto"),
        ("PROMOCAO_APTO", "Promoção — Apto"),
        ("PROMOCAO_INAPTO", "Promoção — Inapto"),
        ("AGREGADO", "Agregado"),
    ]

    status_atual_filtro = [
        ("LTS", "LTS"),
        ("LTSPF", "LTSPF"),
        ("LM", "Licença Maternidade"),
        ("APTO_RECOM", "Apto com Recomendações"),
        ("APTO_RESTR", "Apto com Restrições"),
        ("APTO", "Apto sem Restrição"),
        ("AGUARDANDO_INSPECAO", "Aguardando Inspeção"),
        ("AGREGADO", "Agregado"),
    ]

    return render_template(
        "junta/listar_licencas.html",
        dados=dados_pagina,
        resumo=resumo,
        current_page=page,
        total_pages=total_pages,
        total=total,
        filtro_q=filtros["filtro_q"],
        filtro_tipo=filtros["filtro_tipo"],
        filtro_status=filtros["filtro_status"],
        filtro_status_atual=filtros["filtro_status_atual"],
        filtro_nota_bg=filtros["filtro_nota_bg"],
        filtro_quadro_id=filtros["filtro_quadro_id"],
        filtro_posto_grad_id=filtros["filtro_posto_grad_id"],
        filtro_obm_id=filtros["filtro_obm_id"],
        filtro_restricao_id=filtros["filtro_restricao_id"],
        tipos_filtro=tipos_filtro,
        status_filtro=status_filtro,
        status_atual_filtro=status_atual_filtro,
        quadros=Quadro.query.order_by(Quadro.quadro.asc()).all(),
        postos_grad=PostoGrad.query.order_by(PostoGrad.id.asc()).all(),
        obms=Obm.query.order_by(Obm.sigla.asc()).all(),
        tipos_restricao=listar_tipos_restricao(),
    )


def _ler_filtros_listagem():
    """Filtros da listagem — usados também na exportação e no relatório."""
    return {
        "filtro_q": (request.args.get("q") or "").strip(),
        "filtro_tipo": (request.args.get("tipo") or "").strip(),
        "filtro_status": (request.args.get("status") or "").strip(),
        "filtro_status_atual": (request.args.get("status_atual") or "").strip(),
        "filtro_nota_bg": (request.args.get("nota_bg") or "").strip(),
        "filtro_quadro_id": (request.args.get("quadro_id") or "").strip(),
        "filtro_posto_grad_id": (request.args.get("posto_grad_id") or "").strip(),
        "filtro_obm_id": (request.args.get("obm_id") or "").strip(),
        "filtro_restricao_id": (request.args.get("restricao_id") or "").strip(),
    }


@junta_bp.route("/historico/<int:militar_id>", methods=["GET"])
@login_required
@require_perm("JUNTA_READ")
def historico_militar(militar_id):
    militar = (
        Militar.query
        .options(
            joinedload(Militar.posto_grad),
            joinedload(Militar.quadro),
        )
        .get_or_404(militar_id)
    )

    registros = (
        Licencas.query
        .options(joinedload(Licencas.restricoes).joinedload(LicencaRestricao.tipo))
        .filter_by(militar_id=militar_id)
        .order_by(Licencas.data_inicio.desc(), Licencas.id.desc())
        .all()
    )

    situacao = calcular_situacao_atual(registros)

    return render_template(
        "junta/historico_militar.html",
        militar=militar,
        registros=registros,
        status_atual=situacao["status_atual"],
        status_atual_label=situacao["status_atual_label"],
        agregacao=situacao["agregacao"],
        label_status=label_status,
        label_tipo=label_tipo,
        tipos_restricao=listar_tipos_restricao(),
        tipos_com_restricao=sorted(TIPOS_COM_RESTRICAO),
    )


@junta_bp.route("/api/militares/buscar", methods=["GET"])
@login_required
@require_perm("JUNTA_READ")
def buscar_militares():
    q = (request.args.get("q") or "").strip()

    if len(q) < 1:
        return jsonify([])

    # `_` e `%` são curingas do LIKE: escapa pra que o operador digitando
    # esses caracteres não receba a lista inteira.
    termo = q.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")

    comeca_com = f"{termo}%"
    contem = f"%{termo}%"

    # Ordem alfabética "conforme ele for digitando": primeiro quem COMEÇA com
    # o que foi digitado (nome completo, depois nome de guerra), e só então
    # quem apenas contém o termo no meio — cada bloco em ordem alfabética.
    prioridade = case(
        (Militar.nome_completo.ilike(comeca_com, escape="\\"), 0),
        (Militar.nome_guerra.ilike(comeca_com, escape="\\"), 1),
        else_=2
    )

    militares = (
        Militar.query
        .options(joinedload(Militar.posto_grad))
        .filter(
            Militar.inativo.isnot(True),
            or_(
                Militar.nome_completo.ilike(contem, escape="\\"),
                Militar.nome_guerra.ilike(contem, escape="\\"),
            )
        )
        .order_by(prioridade.asc(), Militar.nome_completo.asc())
        .limit(20)
        .all()
    )

    resultados = []
    for m in militares:
        resultados.append({
            "id": m.id,
            "nome": m.nome_completo or "",
            "nome_guerra": m.nome_guerra or "",
            "matricula": m.matricula or "",
            "posto_grad": m.posto_grad.sigla if m.posto_grad else "",
            "label": f"{m.nome_completo}"
        })

    return jsonify(resultados)


@junta_bp.route("/api/militar/<int:militar_id>", methods=["GET"])
@login_required
@require_perm("JUNTA_READ")
def get_militar_info(militar_id):
    militar = (
        Militar.query
        .options(
            joinedload(Militar.posto_grad),
            joinedload(Militar.quadro),
        )
        .get_or_404(militar_id)
    )

    return jsonify({
        "id": militar.id,
        "nome": militar.nome_completo or "",
        "nome_guerra": militar.nome_guerra or "",
        "matricula": militar.matricula or "",
        "posto_grad": militar.posto_grad.sigla if militar.posto_grad else "N/D",
        "quadro": militar.quadro.quadro if militar.quadro else "N/D",
        "obm": _obter_obm_atual(militar)
    })


@junta_bp.route("/licencas/exportar-excel", methods=["GET"])
@login_required
@require_perm("JUNTA_EXPORT")
def exportar_licencas_excel():
    filtros = _ler_filtros_listagem()
    dados, resumo = montar_dados_licencas(**filtros)

    # Planilha organizada por hierarquia de posto/graduação — do mais antigo
    # (Coronel) pro mais moderno (Aluno-Soldado) — em vez da ordem de
    # lançamento que a tela usa.
    dados.sort(key=lambda item: ordem_hierarquica_militar(item["registro"].militar))

    wb = Workbook()
    ws = wb.active
    ws.title = "Licenças Junta"

    headers = [
        "Militar",
        "Posto/Grad",
        "Quadro",
        "Tipo de Inspeção",
        "Curso",
        "Status do Registro",
        "Status Atual",
        "Restrições",
        "BG",
        "Dias",
        "Data Início",
        "Data Fim",
        "Sessão",
        "Data da Sessão",
        "Agregação",
        "Observação",
        "Criado em",
    ]
    ws.append(headers)

    fill = PatternFill("solid", fgColor="0B2F4F")
    font = Font(color="FFFFFF", bold=True)

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for item in dados:
        reg = item["registro"]
        agregacao = item["agregacao"]
        militar = reg.militar

        if agregacao["atingiu_limite"]:
            agg_texto = "AGREGADO/Apto à agregação"
        elif agregacao["alerta"]:
            agg_texto = agregacao["mensagem"] or "Próximo da agregação"
        else:
            agg_texto = "-"

        ws.append([
            militar.nome_completo if militar else "",
            militar.posto_grad.sigla if militar and militar.posto_grad else "",
            militar.quadro.quadro if militar and militar.quadro else "",
            item["tipo_label"],
            item["curso_nome"] or "",
            item["status_label"],
            item["status_atual_label"],
            "; ".join(item["restricoes"]),
            reg.recebimento_bg,
            reg.qtd_dias,
            reg.data_inicio.strftime("%d/%m/%Y") if reg.data_inicio else "",
            reg.data_fim.strftime("%d/%m/%Y") if reg.data_fim else "",
            reg.sessao,
            reg.data_sessao.strftime("%d/%m/%Y") if reg.data_sessao else "",
            agg_texto,
            reg.observacao or "",
            reg.created_at.strftime(
                "%d/%m/%Y %H:%M") if reg.created_at else "",
        ])

    # Bloco de percentuais sobre o efetivo ativo, logo abaixo da tabela.
    ws.append([])
    ws.append([f"Efetivo ativo considerado: {resumo['efetivo_ativo']}"])
    ws.append(["Situação", "Militares", "% do efetivo ativo"])

    for rotulo, chave_qtd, chave_pct in [
        ("Em licença", "em_licenca", "pct_em_licenca"),
        ("Aptos", "aptos", "pct_aptos"),
        ("Com recomendações", "recomendacoes", "pct_recomendacoes"),
        ("Com restrições", "restricoes", "pct_restricoes"),
        ("Agregados", "agregados", "pct_agregados"),
        ("Aguardando inspeção", "aguardando_inspecao", "pct_aguardando_inspecao"),
    ]:
        ws.append([rotulo, resumo[chave_qtd], f"{resumo[chave_pct]}%"])

    if resumo["restricoes_detalhe"]:
        ws.append([])
        ws.append(["Restrição", "Militares", "% do efetivo ativo"])
        for linha in resumo["restricoes_detalhe"]:
            ws.append([linha["nome"], linha["militares"],
                      f"{linha['percentual']}%"])

    larguras = {
        "A": 42, "B": 14, "C": 14, "D": 26, "E": 28, "F": 30, "G": 30,
        "H": 45, "I": 16, "J": 10, "K": 14, "L": 14, "M": 18, "N": 16,
        "O": 38, "P": 45, "Q": 18
    }
    for col, width in larguras.items():
        ws.column_dimensions[col].width = width

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name="licencas_junta.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


@junta_bp.route("/licencas/relatorio", methods=["GET"])
@login_required
@require_perm("JUNTA_READ")
def relatorio_licencas():
    filtros = _ler_filtros_listagem()
    dados, resumo = montar_dados_licencas(**filtros)

    return render_template(
        "junta/relatorio_licencas.html",
        dados=dados,
        resumo=resumo,
        hoje=hoje_manaus(),
        filtro_q=filtros["filtro_q"],
        filtro_tipo=filtros["filtro_tipo"],
        filtro_status=filtros["filtro_status"],
        filtro_status_atual=filtros["filtro_status_atual"],
    )


@junta_bp.route("/licencas/finalizar-bg", methods=["POST"])
@login_required
@require_perm("JUNTA_BG_FECHAR")
def finalizar_bg_dia():
    eh_errata = request.form.get("eh_errata") == "on"

    if eh_errata:
        return _finalizar_errata()
    return _finalizar_fechamento_dia()


def _finalizar_fechamento_dia():
    """
    Fechamento normal: agrupa os lançamentos pendentes de UMA sessão da Junta
    (sessao + data_sessao) numa nota. Não usa a data em que foi digitado no
    sistema — usa a sessão de verdade, escolhida entre as que ainda têm
    lançamento pendente. É assim que dá pra fazer quantas notas quiser, em
    dias e sessões diferentes, sem misturar uma coisa com a outra.
    """
    fechamento = None

    try:
        chave = (request.form.get("sessao_data_pendente") or "").strip()
        nota_bg = (request.form.get("nota_bg") or "").strip()
        observacao_bg = (request.form.get("observacao_bg") or "").strip()

        if not chave or "||" not in chave:
            flash("Selecione a sessão que será fechada.", "danger")
            return redirect(url_for("junta.nova_licenca"))

        sessao_pendente, data_pendente_str = chave.split("||", 1)

        if not nota_bg:
            flash("Informe a nota para BG.", "danger")
            return redirect(url_for("junta.nova_licenca"))

        try:
            data_pendente = datetime.strptime(
                data_pendente_str, "%Y-%m-%d").date()
        except ValueError:
            flash("Sessão selecionada é inválida.", "danger")
            return redirect(url_for("junta.nova_licenca"))

        pendentes = (
            Licencas.query
            .filter(
                Licencas.sessao == sessao_pendente,
                Licencas.data_sessao == data_pendente,
                Licencas.fechamento_bg_id.is_(None)
            )
            .order_by(Licencas.created_at.asc(), Licencas.id.asc())
            .all()
        )

        if not pendentes:
            flash(
                "Não há mais lançamentos pendentes para essa sessão — "
                "talvez ela já tenha sido fechada. Atualize a página.",
                "warning"
            )
            return redirect(url_for("junta.nova_licenca"))

        fechamento = JuntaFechamentoBg(
            data_referencia=data_pendente,
            nota_bg=nota_bg,
            sessao=sessao_pendente,
            observacao=observacao_bg or None,
            usuario_id=current_user.id
        )

        database.session.add(fechamento)
        database.session.flush()

        for lic in pendentes:
            lic.fechamento_bg_id = fechamento.id

        # ainda sem commit final
        arquivo = gerar_nota_bg_docx(fechamento.id, commit_db=False)

        fechamento.arquivo_docx = arquivo
        database.session.commit()

        flash(
            f"Fechamento realizado com sucesso. {len(pendentes)} lançamento(s) "
            f"da sessão {sessao_pendente} vinculados à nota BG {nota_bg}.",
            "success"
        )

    except Exception as e:
        database.session.rollback()
        flash(f"Erro ao finalizar BG do dia: {str(e)}", "danger")
        return redirect(url_for("junta.nova_licenca"))

    return redirect(
        url_for(
            "junta.baixar_docx_fechamento",
            fechamento_id=fechamento.id
        )
    )


def _finalizar_errata():
    """
    Errata: corrige uma nota já publicada. Não mexe em lançamentos pendentes
    — é um documento à parte, referenciando a nota original e o Boletim Geral
    onde ela de fato saiu (que nem sempre é do mesmo dia da sessão).
    """
    fechamento = None

    try:
        nota_bg = (request.form.get("nota_bg") or "").strip()
        sessao_bg = (request.form.get("sessao_bg") or "").strip()
        fechamento_original_id = request.form.get(
            "fechamento_original_id", type=int)
        numero_bg_publicacao = (
            request.form.get("numero_bg_publicacao") or "").strip()
        data_bg_publicacao_str = (
            request.form.get("data_bg_publicacao") or "").strip()
        onde_se_le = (request.form.get("onde_se_le") or "").strip()
        leia_se = (request.form.get("leia_se") or "").strip()
        observacao_bg = (request.form.get("observacao_bg") or "").strip()

        if not nota_bg:
            flash("Informe a nota para BG da errata.", "danger")
            return redirect(url_for("junta.nova_licenca"))

        if not sessao_bg:
            flash("Informe a sessão.", "danger")
            return redirect(url_for("junta.nova_licenca"))

        if not fechamento_original_id:
            flash("Selecione a nota que está sendo corrigida.", "danger")
            return redirect(url_for("junta.nova_licenca"))

        original = JuntaFechamentoBg.query.get(fechamento_original_id)
        if not original:
            flash("Nota original não encontrada.", "danger")
            return redirect(url_for("junta.nova_licenca"))

        if not numero_bg_publicacao:
            flash(
                "Informe o número do Boletim Geral onde a nota original foi publicada.", "danger")
            return redirect(url_for("junta.nova_licenca"))

        if not data_bg_publicacao_str:
            flash("Informe a data do Boletim Geral.", "danger")
            return redirect(url_for("junta.nova_licenca"))

        if not onde_se_le or not leia_se:
            flash(
                "Preencha o que estava escrito (\"onde se lê\") e a correção (\"leia-se\").",
                "danger"
            )
            return redirect(url_for("junta.nova_licenca"))

        data_bg_publicacao = datetime.strptime(
            data_bg_publicacao_str, "%Y-%m-%d").date()

        fechamento = JuntaFechamentoBg(
            data_referencia=hoje_manaus(),
            nota_bg=nota_bg,
            sessao=sessao_bg,
            observacao=observacao_bg or None,
            usuario_id=current_user.id,
            eh_errata=True,
            fechamento_original_id=original.id,
            numero_bg_publicacao=numero_bg_publicacao,
            data_bg_publicacao=data_bg_publicacao,
            onde_se_le=onde_se_le,
            leia_se=leia_se,
        )

        database.session.add(fechamento)
        database.session.flush()

        arquivo = gerar_nota_bg_docx(fechamento.id, commit_db=False)
        fechamento.arquivo_docx = arquivo
        database.session.commit()

        flash(
            f"Errata da nota BG {original.nota_bg} registrada com sucesso.",
            "success"
        )

    except Exception as e:
        database.session.rollback()
        flash(f"Erro ao registrar errata: {str(e)}", "danger")
        return redirect(url_for("junta.nova_licenca"))

    return redirect(
        url_for("junta.baixar_docx_fechamento", fechamento_id=fechamento.id)
    )


@junta_bp.route("/fechamento-bg/<int:fechamento_id>/baixar-docx", methods=["GET"])
@login_required
@require_perm("JUNTA_BG_FECHAR")
def baixar_docx_fechamento(fechamento_id):
    fechamento = JuntaFechamentoBg.query.get_or_404(fechamento_id)

    if not fechamento.arquivo_docx:
        flash("Este fechamento ainda não possui documento gerado.", "warning")
        return redirect(url_for("junta.listar_licencas"))

    pasta = Path("src/static/junta_bg")
    return send_from_directory(
        directory=str(pasta.resolve()),
        path=fechamento.arquivo_docx,
        as_attachment=True
    )


def _ler_filtros_estatisticas():
    """Período e recortes do painel mensal."""
    hoje = hoje_manaus()

    # Padrão: o ano corrente até o mês atual.
    padrao_inicio = date(hoje.year, 1, 1)

    inicio = normalizar_competencia(request.args.get("de"), padrao_inicio)
    fim = normalizar_competencia(request.args.get("ate"), hoje)

    if fim < inicio:
        inicio, fim = fim, inicio

    return {
        "inicio": inicio,
        "fim": fim,
        "filtro_quadro_id": (request.args.get("quadro_id") or "").strip(),
        "filtro_posto_grad_id": (request.args.get("posto_grad_id") or "").strip(),
        "filtro_obm_id": (request.args.get("obm_id") or "").strip(),
        "filtro_tipo": (request.args.get("tipo") or "").strip(),
    }


@junta_bp.route("/estatisticas", methods=["GET"])
@login_required
@require_perm("JUNTA_READ")
def estatisticas_mensais():
    """
    Painel mensal. Sem `militar_id` mostra o panorama da força; com
    `militar_id` mostra o mesmo recorte para um militar só.
    """
    filtros = _ler_filtros_estatisticas()
    militar_id = request.args.get("militar_id", type=int)

    militar = None
    situacao = None

    if militar_id:
        militar = (
            Militar.query
            .options(
                joinedload(Militar.posto_grad),
                joinedload(Militar.quadro),
            )
            .get_or_404(militar_id)
        )

        historico = (
            Licencas.query
            .filter_by(militar_id=militar.id)
            .order_by(Licencas.data_inicio.desc(), Licencas.id.desc())
            .all()
        )
        situacao = calcular_situacao_atual(historico)

    resultado = montar_estatisticas_mensais(
        filtros["inicio"],
        filtros["fim"],
        filtro_quadro_id=filtros["filtro_quadro_id"],
        filtro_posto_grad_id=filtros["filtro_posto_grad_id"],
        filtro_obm_id=filtros["filtro_obm_id"],
        filtro_tipo=filtros["filtro_tipo"],
        militar_id=militar_id,
    )

    # O ranking só faz sentido no panorama geral — com um militar escolhido,
    # o lugar dele é a própria tabela mensal.
    ranking = [] if militar_id else montar_ranking_militares(resultado)

    efetivo_ativo = contar_efetivo_ativo(
        filtro_quadro_id=filtros["filtro_quadro_id"],
        filtro_posto_grad_id=filtros["filtro_posto_grad_id"],
        filtro_obm_id=filtros["filtro_obm_id"],
    )

    tipos_filtro = [
        ("LTS", "LTS"),
        ("LTSPF", "LTSPF"),
        ("LM", "Licença Maternidade"),
        ("APTO_RECOM", "Apto com Recomendações"),
        ("APTO_RESTR", "Apto com Restrições"),
        ("APTO", "Apto sem Restrição"),
        ("CURSO", "Curso"),
        ("TAF", "TAF"),
        ("PROMOCAO", "Promoção"),
        ("AGREGADO", "Agregado"),
    ]

    return render_template(
        "junta/estatisticas.html",
        linhas=resultado["linhas"],
        totais=resultado["totais"],
        tipos_presentes=resultado["tipos"],
        restricoes_presentes=resultado["restricoes"],
        registros=resultado["registros"],
        grafico=resultado["grafico"],
        ranking=ranking,
        militar=militar,
        situacao=situacao,
        efetivo_ativo=efetivo_ativo,
        filtro_de=filtros["inicio"].strftime("%Y-%m"),
        filtro_ate=filtros["fim"].strftime("%Y-%m"),
        filtro_quadro_id=filtros["filtro_quadro_id"],
        filtro_posto_grad_id=filtros["filtro_posto_grad_id"],
        filtro_obm_id=filtros["filtro_obm_id"],
        filtro_tipo=filtros["filtro_tipo"],
        tipos_filtro=tipos_filtro,
        quadros=Quadro.query.order_by(Quadro.quadro.asc()).all(),
        postos_grad=PostoGrad.query.order_by(PostoGrad.id.asc()).all(),
        obms=Obm.query.order_by(Obm.sigla.asc()).all(),
        label_tipo=label_tipo,
    )


@junta_bp.route("/estatisticas/exportar-excel", methods=["GET"])
@login_required
@require_perm("JUNTA_EXPORT")
def exportar_estatisticas_excel():
    filtros = _ler_filtros_estatisticas()
    militar_id = request.args.get("militar_id", type=int)

    # get_or_404 igual à tela: um militar_id inválido tem que dar 404, e não
    # baixar uma planilha vazia sem dizer o motivo.
    militar = Militar.query.get_or_404(militar_id) if militar_id else None

    resultado = montar_estatisticas_mensais(
        filtros["inicio"],
        filtros["fim"],
        filtro_quadro_id=filtros["filtro_quadro_id"],
        filtro_posto_grad_id=filtros["filtro_posto_grad_id"],
        filtro_obm_id=filtros["filtro_obm_id"],
        filtro_tipo=filtros["filtro_tipo"],
        militar_id=militar_id,
    )

    linhas = resultado["linhas"]
    totais = resultado["totais"]
    tipos = resultado["tipos"]

    wb = Workbook()
    ws = wb.active
    ws.title = "Consolidado mensal"

    fill = PatternFill("solid", fgColor="0B2F4F")
    fonte = Font(color="FFFFFF", bold=True)
    fonte_total = Font(bold=True)

    if militar:
        ws.append([f"Militar: {militar.nome_completo}"])
        ws.append([])

    ws.append([
        f"Período: {filtros['inicio'].strftime('%m/%Y')} a "
        f"{filtros['fim'].strftime('%m/%Y')}"
    ])
    ws.append([])

    # Mesma troca que a tela faz: com um militar escolhido, "militares
    # distintos" seria sempre 0 ou 1 — a coluna vira a contagem de pareceres.
    cabecalho = [
        "Competência",
        "Licenças lançadas",
        "Restrições lançadas",
        "Inspeções lançadas",
        "Licenças vigentes",
        "Dias de licença",
        "Pareceres" if militar else "Militares distintos",
    ] + [label for _, label in tipos]

    ws.append(cabecalho)
    linha_cabecalho = ws.max_row

    for col in range(1, len(cabecalho) + 1):
        celula = ws.cell(row=linha_cabecalho, column=col)
        celula.fill = fill
        celula.font = fonte
        celula.alignment = Alignment(horizontal="center", vertical="center")

    for linha in linhas:
        ws.append([
            linha["competencia"],
            linha["licencas_lancadas"],
            linha["restricoes_lancadas"],
            linha["inspecoes_lancadas"],
            linha["licencas_vigentes"],
            linha["dias_licenca"],
            linha["inspecoes_lancadas"] if militar else linha["militares_distintos"],
        ] + [linha["por_tipo"].get(tipo, 0) for tipo, _ in tipos])

    ws.append([
        "TOTAL",
        totais["licencas_lancadas"],
        totais["restricoes_lancadas"],
        totais["inspecoes_lancadas"],
        "-",                      # vigentes não soma: o mesmo registro
                                  # atravessa vários meses
        totais["dias_licenca"],
        totais["inspecoes_lancadas"] if militar else totais["militares_distintos"],
    ] + [totais["por_tipo"].get(tipo, 0) for tipo, _ in tipos])

    for col in range(1, len(cabecalho) + 1):
        ws.cell(row=ws.max_row, column=col).font = fonte_total

    ws.column_dimensions["A"].width = 16
    for col in "BCDEFG":
        ws.column_dimensions[col].width = 20

    # --- aba de restrições por mês ---
    if resultado["restricoes"]:
        ws2 = wb.create_sheet("Restrições por mês")
        cab2 = ["Competência"] + resultado["restricoes"] + ["Total do mês"]
        ws2.append(cab2)

        for col in range(1, len(cab2) + 1):
            celula = ws2.cell(row=1, column=col)
            celula.fill = fill
            celula.font = fonte
            celula.alignment = Alignment(horizontal="center", vertical="center")

        for linha in linhas:
            ws2.append(
                [linha["competencia"]]
                + [linha["por_restricao"].get(nome, 0)
                   for nome in resultado["restricoes"]]
                + [linha["restricoes_lancadas"]]
            )

        ws2.append(
            ["TOTAL"]
            + [totais["por_restricao"].get(nome, 0)
               for nome in resultado["restricoes"]]
            + [totais["restricoes_lancadas"]]
        )
        for col in range(1, len(cab2) + 1):
            ws2.cell(row=ws2.max_row, column=col).font = fonte_total

        ws2.column_dimensions["A"].width = 16
        for i in range(2, len(cab2) + 1):
            ws2.column_dimensions[ws2.cell(row=1, column=i).column_letter].width = 32

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    sufixo = f"_militar_{militar_id}" if militar_id else ""
    return send_file(
        output,
        as_attachment=True,
        download_name=f"junta_consolidado_mensal{sufixo}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


@junta_bp.route("/renovacoes", methods=["GET"])
@login_required
@require_perm("JUNTA_RENOVACOES_READ")
def painel_renovacoes():
    mes = request.args.get("mes", type=int)
    ano = request.args.get("ano", type=int)
    tipo = (request.args.get("tipo") or "").strip()

    query = (
        Licencas.query
        .options(
            joinedload(Licencas.militar).joinedload(Militar.posto_grad),
            joinedload(Licencas.militar).joinedload(Militar.quadro),
        )
        .order_by(Licencas.militar_id.asc(), Licencas.data_inicio.asc(), Licencas.id.asc())
    )

    if tipo:
        query = query.filter(Licencas.tipo_licenca == tipo)

    if mes and ano:
        inicio_mes = date(ano, mes, 1)
        fim_mes = date(ano, mes, monthrange(ano, mes)[1])

        query = query.filter(
            Licencas.data_inicio <= fim_mes,
            Licencas.data_fim >= inicio_mes
        )

    registros = query.all()
    blocos_por_militar = montar_blocos_por_militar(registros)

    linhas = []

    for militar_id, blocos in blocos_por_militar.items():
        for bloco in blocos:
            primeiro_reg = bloco["registros"][0]
            militar = primeiro_reg.militar

            linhas.append({
                "militar": militar,
                "tipo_licenca": bloco["tipo_licenca"],
                "tipo_label": bloco["tipo_label"],
                "inicio_bloco": bloco["inicio_bloco"],
                "fim_bloco": bloco["fim_bloco"],
                "dias_continuos": bloco["dias_continuos"],
                "renovacoes": bloco["renovacoes"],
                "quantidade_registros": bloco["quantidade_registros"],
                "ultima_renovacao": bloco["ultima_renovacao"],
                "meses_abrangidos": bloco["meses_abrangidos"],
                "datas_renovacao": bloco["datas_renovacao"],
                "suspeito": bloco["suspeito"],
                "motivo_suspeita": bloco["motivo_suspeita"],
            })

    linhas.sort(key=lambda x: (
        x["suspeito"], x["renovacoes"], x["dias_continuos"]), reverse=True)

    resumo = {
        "total_blocos": len(linhas),
        "com_renovacao": sum(1 for x in linhas if x["renovacoes"] > 0),
        "suspeitos": sum(1 for x in linhas if x["suspeito"]),
    }

    return render_template(
        "junta/painel_renovacoes.html",
        linhas=linhas,
        resumo=resumo,
        filtro_mes=mes,
        filtro_ano=ano,
        filtro_tipo=tipo,
    )
